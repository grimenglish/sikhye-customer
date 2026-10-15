
import io
import re
import json
import math
import hmac
import time
from datetime import datetime, date

from datetime import timedelta
from zoneinfo import ZoneInfo


import pandas as pd
import streamlit as st
import requests
import altair as alt

try:
    import msoffcrypto
except Exception:
    msoffcrypto = None

try:
    from supabase import create_client
except Exception:
    create_client = None


st.set_page_config(
    page_title="식혜명가 저장형 고객 CRM",
    page_icon="📦",
    layout="wide",
)

st.markdown("""
<style>
.block-container {padding-top: 1.4rem; max-width: 1500px; padding-bottom: 3rem;}
.kpi-card {
    background: white;
    border: 1px solid #ECECEC;
    border-radius: 18px;
    padding: 20px 22px;
    box-shadow: 0 2px 8px rgba(0,0,0,0.04);
}
.kpi-title {color:#666;font-size:14px;margin-bottom:8px;}
.kpi-value {font-size:28px;font-weight:800;color:#111;}
.note {
    background:#f6f8fa;
    border:1px solid #e5e7eb;
    padding:14px 16px;
    border-radius:12px;
}
.warn {
    background:#fff7ed;
    border:1px solid #fed7aa;
    padding:14px 16px;
    border-radius:12px;
}

.upload-guide {
    border: 2px dashed #cbd5e1;
    background: #f8fafc;
    border-radius: 18px;
    padding: 22px;
    text-align: center;
    margin-bottom: 14px;
}
.upload-title {
    font-size: 22px;
    font-weight: 800;
    margin-bottom: 6px;
}
.upload-sub {
    color: #64748b;
    font-size: 14px;
}
.summary-box {
    background: #ecfdf5;
    border: 1px solid #bbf7d0;
    border-radius: 16px;
    padding: 16px 18px;
    margin: 12px 0;
}


/* V12 시각 정리: 기존 기능·위젯 구조는 유지 */
:root { color-scheme: light; }
.stTabs [data-baseweb="tab-list"] {
    gap: 4px;
    overflow-x: auto;
    scrollbar-width: thin;
    border-bottom: 1px solid #e5e7eb;
}
.stTabs [data-baseweb="tab"] {
    white-space: nowrap;
    border-radius: 10px 10px 0 0;
    padding: 10px 12px;
}
.stTabs [aria-selected="true"] {
    color: #b45309 !important;
    background: #fff7ed !important;
    font-weight: 700 !important;
}
[data-testid="stMetric"] {
    border: 1px solid #e5e7eb;
    border-radius: 12px;
    padding: 12px 14px;
    background-color: #ffffff;
}
[data-testid="stVerticalBlockBorderWrapper"]:has([data-testid="stNumberInput"]) {
    border-color: #e5e7eb;
}
@media (max-width: 650px) {
    .block-container {padding-top: 1rem; padding-left: .65rem; padding-right: .65rem;}
    h1 {font-size: 1.7rem !important;}
    h2 {font-size: 1.35rem !important;}
    [data-testid="stMetric"] {padding: 8px;}
    .stTabs [data-baseweb="tab"] {padding: 8px 10px;}
}

/* 모바일/다크모드 표 가독성 보정 */
[data-testid="stDataFrame"] {
    background: #ffffff !important;
}
[data-testid="stDataFrame"] div {
    font-size: 15px;
}


/* V21: 실제 업로더를 3배 가까이 크게, 버튼은 터치하기 편하게 */
[data-testid="stFileUploaderDropzone"] {
    min-height: 345px !important;
    padding: 65px 28px !important;
    border: 3px dashed #94a3b8 !important;
    border-radius: 22px !important;
    background: linear-gradient(155deg, #f8fbff, #eef5ff) !important;
    cursor: pointer !important;
    transition: background .2s ease, border-color .2s ease !important;
}
[data-testid="stFileUploaderDropzone"]:hover {
    border-color: #2563eb !important;
    background: #eaf3ff !important;
}
[data-testid="stFileUploaderDropzone"] button {
    min-height: 58px !important;
    min-width: 165px !important;
    padding: 12px 24px !important;
    border-radius: 14px !important;
    font-size: 19px !important;
    font-weight: 750 !important;
    color: #123b78 !important;
    border: 2px solid #93b4e8 !important;
    background: white !important;
}
[data-testid="stFileUploaderDropzoneInstructions"] {
    min-width: 0 !important;
    text-align: center !important;
}
[data-testid="stFileUploaderDropzoneInstructions"] * {
    font-size: 18px !important;
    line-height: 1.55 !important;
}
[data-testid="stFileUploaderDropzone"] svg { width: 42px !important; height: 42px !important; }
[data-testid="stFileUploader"] small { font-size: 14px !important; }
[data-testid="stFileUploader"] label { font-size: 17px !important; font-weight: 750 !important; }

/* CRM 전체 텍스트 가독성 + 버튼 클릭 면적 */
[data-testid="stMetricLabel"] p { font-size: 15px !important; }
[data-testid="stMetricValue"] div { font-size: clamp(22px, 2vw, 32px) !important; font-weight: 750 !important; }
.stTabs [data-baseweb="tab"] { font-size: 15px !important; min-height: 47px !important; }
.stButton button[kind="primary"] { min-height: 45px; border-radius: 11px; font-weight: 700; }
[data-testid="stDataFrame"] { border-radius: 10px; overflow: hidden; }
.vg-tooltip, #vg-tooltip-element { font-size: 17px !important; line-height: 1.55 !important; }
.vg-tooltip td, .vg-tooltip th { font-size: 16px !important; padding: 7px 11px !important; }
@media (max-width: 650px) {
    [data-testid="stFileUploaderDropzone"] { min-height: 245px !important; padding: 33px 12px !important; }
    [data-testid="stFileUploaderDropzone"] button { min-height: 48px !important; min-width: 130px !important; font-size: 16px !important; }
    [data-testid="stFileUploaderDropzoneInstructions"] * { font-size: 15px !important; }
}

/* V22: 그래프 축/수치 선명도 및 경영 대시보드 간격 */
[data-testid="stMetric"] [data-testid="stMetricDelta"] { font-size: 15px !important; }
[data-testid="stMetric"] [data-testid="stMetricLabel"] { line-height: 1.45 !important; }
[data-testid="stAltairChart"] { min-width: 0; }
[data-testid="stAltairChart"] canvas { max-width: 100% !important; }
@media (max-width: 680px) {
  [data-testid="stMetric"] [data-testid="stMetricValue"] { font-size: 25px !important; }
  [data-testid="stMetric"] [data-testid="stMetricLabel"] p { font-size: 14px !important; }
}

</style>
""", unsafe_allow_html=True)

SMARTSTORE_PASSWORD = "1111"



# =========================
# 앱 비밀번호 로그인
# =========================
def require_login():
    """공유 비밀번호 접근 보호. 5회 실패 시 세션별 일시 잠금, 장시간 미사용 시 재로그인."""
    app_password = str(st.secrets.get("APP_PASSWORD", ""))
    if not app_password:
        st.error("APP_PASSWORD가 Streamlit Secrets에 설정되지 않았습니다. 개인정보 보호를 위해 앱을 열 수 없습니다.")
        st.stop()

    now = time.monotonic()
    locked_until = float(st.session_state.get("_login_locked_until", 0.0))
    last_activity = st.session_state.get("_login_last_activity")
    if st.session_state.get("authenticated") and last_activity is not None and now - last_activity > 3 * 3600:
        st.session_state.clear()
        st.warning("3시간 이상 사용 기록이 없어 로그아웃됐습니다. 다시 로그인해 주세요.")
    if st.session_state.get("authenticated"):
        st.session_state["_login_last_activity"] = now
        with st.sidebar:
            st.success("로그인됨")
            if st.button("로그아웃"):
                st.session_state.clear()
                st.rerun()
        return

    st.title("🔐 식혜명가 CRM 로그인")
    st.caption("고객 개인정보 보호를 위해 비밀번호가 필요합니다.")
    if locked_until > now:
        remaining = int((locked_until - now + 59) // 60)
        st.error(f"로그인 시도가 많아 잠시 제한됐습니다. 약 {remaining}분 후 다시 시도해 주세요.")
        st.stop()

    password = st.text_input("비밀번호", type="password")
    if st.button("로그인", type="primary"):
        if hmac.compare_digest(str(password), app_password):
            st.session_state["authenticated"] = True
            st.session_state["_login_last_activity"] = now
            st.session_state.pop("_login_failed", None)
            st.session_state.pop("_login_locked_until", None)
            st.rerun()
        else:
            failed = int(st.session_state.get("_login_failed", 0)) + 1
            st.session_state["_login_failed"] = failed
            if failed >= 5:
                st.session_state["_login_locked_until"] = now + 15 * 60
                st.session_state["_login_failed"] = 0
                st.error("비밀번호 오류가 반복되어 15분간 로그인이 제한됩니다.")
            else:
                st.error("비밀번호가 틀렸습니다.")
    st.stop()


require_login()


# =========================
# Supabase 연결
# =========================
@st.cache_resource(show_spinner=False)
def _cached_supabase_client(url, key):
    return create_client(url, key)


def get_supabase():
    if create_client is None:
        return None
    url = st.secrets.get("SUPABASE_URL", "")
    key = st.secrets.get("SUPABASE_KEY", "")
    if not url or not key:
        return None
    return _cached_supabase_client(url, key)


sb = get_supabase()


# =========================
# 유틸
# =========================
def normalize_phone(x):
    if pd.isna(x):
        return ""
    return re.sub(r"[^0-9]", "", str(x))


def normalize_text(x):
    if pd.isna(x):
        return ""
    return re.sub(r"\s+", " ", str(x)).strip()


def normalize_address(x):
    s = normalize_text(x)
    s = re.sub(r"\([^)]*\)", "", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def pick_col(df, candidates):
    cols = list(df.columns)
    exact = {str(c).strip(): c for c in cols}

    for c in candidates:
        if c in exact:
            return exact[c]

    for col in cols:
        col_s = str(col).strip()
        for c in candidates:
            if c in col_s:
                return col

    return None


def safe_series(df, col, default=""):
    if col and col in df.columns:
        return df[col]
    return pd.Series([default] * len(df), index=df.index)


def read_excel_file(uploaded_file):
    """일반 xlsx, 실제 xls, 암호화된 네이버 xlsx를 구분해 읽는다.

    암호가 없는 .xls를 복호화하려 하지 않는다.
    """
    raw = uploaded_file.getvalue()
    name = str(getattr(uploaded_file, "name", "")).lower()
    is_ole = raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
    is_zip = raw.startswith(b"PK")
    if not (is_ole or is_zip):
        raise ValueError("엑셀 파일 형식이 아닙니다. 실제 .xls 또는 .xlsx 파일을 선택해 주세요.")

    if is_zip:
        try:
            return pd.read_excel(io.BytesIO(raw), engine="openpyxl")
        except Exception as exc:
            raise ValueError(f"XLSX 파일 읽기에 실패했습니다: {exc}") from exc

    # OLE 파일은 비암호화 xls일 수도 있고 암호화된 xlsx일 수도 있다.
    # 암호화 여부를 확인하지 않고 OfficeFile.decrypt를 실행하면 오해의 소지가 있다.
    if msoffcrypto is not None:
        try:
            office = msoffcrypto.OfficeFile(io.BytesIO(raw))
            encrypted = office.is_encrypted()
        except Exception:
            encrypted = False
        if encrypted:
            try:
                office.load_key(password=str(st.secrets.get("SMARTSTORE_PASSWORD", SMARTSTORE_PASSWORD)))
                decrypted = io.BytesIO()
                office.decrypt(decrypted)
                decrypted.seek(0)
                return pd.read_excel(decrypted, engine="openpyxl")
            except Exception as exc:
                raise ValueError("암호화 엑셀 해제에 실패했습니다. 네이버 원본 파일 및 암호 설정을 확인해 주세요.") from exc

    try:
        return pd.read_excel(io.BytesIO(raw), engine="xlrd")
    except ImportError as exc:
        raise RuntimeError("XLS(.xls) 읽기 라이브러리 xlrd가 없습니다. requirements.txt에 xlrd==2.0.2를 추가해 주세요.") from exc
    except Exception as exc:
        raise ValueError(f"XLS 파일을 읽을 수 없습니다. 파일이 정상적인 엑셀인지 확인해 주세요: {exc}") from exc



def convert_naver_cj_excel(uploaded_bytes: bytes) -> tuple[bytes, int, int]:
    """네이버의 암호화 발주발송 엑셀에서 1행 안내문만 안전하게 제거한다.

    반환: (암호 없는 XLSX 바이트, 상품행 수, 열 수). DB나 서버 디스크에 저장하지 않는다.
    """
    import zipfile
    from openpyxl import load_workbook

    if not uploaded_bytes or len(uploaded_bytes) > 25 * 1024 * 1024:
        raise ValueError("파일이 비어 있거나 25MB를 초과합니다.")

    ole_magic = bytes.fromhex("D0CF11E0A1B11AE1")
    if uploaded_bytes.startswith(ole_magic):
        if msoffcrypto is None:
            raise RuntimeError("암호화 엑셀 변환에 msoffcrypto-tool 라이브러리가 필요합니다. requirements.txt를 확인하세요.")
        try:
            encrypted = msoffcrypto.OfficeFile(io.BytesIO(uploaded_bytes))
            encrypted.load_key(password=str(st.secrets.get("SMARTSTORE_PASSWORD", SMARTSTORE_PASSWORD)))
            decrypted = io.BytesIO()
            encrypted.decrypt(decrypted)
            payload = decrypted.getvalue()
        except Exception as exc:
            raise ValueError("엑셀 암호를 풀 수 없습니다. 스마트스토어 파일 또는 비밀번호 설정을 확인하세요.") from exc
    elif uploaded_bytes.startswith(b"PK"):
        payload = uploaded_bytes
    else:
        raise ValueError("지원하지 않는 파일 형식입니다. 네이버에서 내려받은 .xlsx 파일을 넣어주세요.")

    if not zipfile.is_zipfile(io.BytesIO(payload)):
        raise ValueError("암호 해제 후 엑셀 형식을 확인할 수 없습니다.")

    workbook = load_workbook(io.BytesIO(payload), data_only=False, keep_links=True)
    try:
        sheet = workbook.worksheets[0]
        first = str(sheet.cell(1, 1).value or "")
        second = [str(sheet.cell(2, i).value or "").strip() for i in range(1, 11)]
        if second[:2] != ["상품주문번호", "주문번호"] or "배송방법" not in second or "송장번호" not in second:
            if str(sheet.cell(1, 1).value or "").strip() == "상품주문번호":
                raise ValueError("이미 안내문 1행이 제거된 파일입니다. 다시 변환할 필요가 없습니다.")
            raise ValueError("파일의 2행이 네이버 발주·발송 양식이 아닙니다. 원본을 확인해 주세요.")
        if "다운로드 받은 파일" not in first and "엑셀 일괄발송" not in first:
            raise ValueError("첫 번째 행이 네이버 안내문인지 확인할 수 없어 삭제하지 않았습니다.")
        if len(workbook.worksheets) > 1:
            raise ValueError("시트가 여러 개라 자동 삭제하지 않았습니다. 파일을 확인해 주세요.")

        # 1행의 병합 셀(A1:C1 등)을 먼저 해제하지 않으면 새 표 헤더가 합쳐질 수 있다.
        for span in list(sheet.merged_cells.ranges):
            if span.min_row == 1:
                if span.max_row != 1:
                    raise ValueError("첫 행과 데이터가 함께 병합되어 있어 자동 변환할 수 없습니다.")
                sheet.unmerge_cells(str(span))

        expected_header = tuple(sheet.cell(2, j).value for j in range(1, sheet.max_column + 1))
        first_data = tuple(sheet.cell(3, j).value for j in range(1, sheet.max_column + 1)) if sheet.max_row > 2 else None
        total_columns = sheet.max_column
        original_rows = sheet.max_row
        sheet.delete_rows(1, 1)
        output = io.BytesIO()
        workbook.save(output)
        result = output.getvalue()
    finally:
        workbook.close()

    # 저장된 파일까지 다시 열어 헤더와 첫 주문행을 비교한다.
    verification = load_workbook(io.BytesIO(result), read_only=True, data_only=False)
    try:
        check = verification.worksheets[0]
        if tuple(check.cell(1, j).value for j in range(1, total_columns + 1)) != expected_header:
            raise RuntimeError("변환된 파일의 제목 행이 원본과 달라 저장을 취소했습니다.")
        if first_data is not None and tuple(check.cell(2, j).value for j in range(1, total_columns + 1)) != first_data:
            raise RuntimeError("변환된 파일의 주문 내용이 변경되어 저장을 취소했습니다.")
        if check.max_row != original_rows - 1:
            raise RuntimeError("변환된 파일의 행 수가 일치하지 않습니다.")
    finally:
        verification.close()
    return result, max(0, original_rows - 2), total_columns

# =========================
# V26: 네이버/CJ 변환 및 4개 파일 자동 매칭 (단일 app.py 내부 함수)
# =========================
def _ship_norm(value):
    """고객 식별용 최소한의 표기 정규화. 이름/주소의 추측 매칭은 하지 않는다."""
    import unicodedata
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).lower()
    return "".join(ch for ch in text if ch.isalnum())


def _ship_digits(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\D", "", str(value))


def _ship_str(value):
    return "" if value is None else str(value).strip()


def _ship_identity(entry):
    return (entry["name"], entry["phone"], entry["address"])


def _ship_signature(entry):
    return (entry["item"], entry["quantity"])


def _ship_extract_records(sheet, header, role):
    """필수 열 이름으로 추출해, 엑셀 열 순서가 바뀌어도 안전하게 처리한다."""
    hd = {str(value or "").strip(): i for i, value in enumerate(header)}
    if role == "naver_order":
        columns = ("수취인명", "수취인연락처1", "통합배송지", "옵션정보", "수량", "송장번호")
    elif role == "coupang_order":
        columns = ("수취인이름", "수취인전화번호", "수취인 주소", "등록옵션명", "구매수(수량)", "운송장번호")
    else:
        columns = ("받는분성명", "받는분전화번호", "받는분주소(전체, 분할)", "내품명", "내품수량", "운송장번호")
    missing = [name for name in columns if name not in hd]
    if missing:
        raise ValueError(f"양식에 필요한 열이 없습니다: {', '.join(missing)}")
    indices = [hd[x] for x in columns]
    result = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = [(cells[k] if k < len(cells) else None) for k in indices]
        if all(x is None or str(x).strip() == "" for x in values):
            continue
        name, phone, address, item, quantity, tracking = values
        row = {
            "row": row_number,
            "name": _ship_norm(name),
            "phone": _ship_digits(phone),
            "address": _ship_norm(address),
            "item": _ship_norm(item),
            "quantity": _ship_str(quantity),
            "tracking": _ship_str(tracking),
            "tracking_digits": _ship_digits(tracking),
        }
        # 잘못된 송장 결합보다 작업 중단이 안전함.
        if not all([row["name"], row["phone"], row["address"], row["item"], row["quantity"]]):
            raise ValueError(f"{role} 파일의 {row_number}행에 이름·전화·주소·옵션·수량 누락이 있습니다.")
        if role == "cj" and (not (10 <= len(row["tracking_digits"]) <= 16)
                             or re.fullmatch(r"[0-9\- ]+", row["tracking"]) is None):
            raise ValueError(f"CJ 파일 {row_number}행의 송장번호가 올바르지 않습니다.")
        if role != "cj" and row["tracking"]:
            raise ValueError(f"주문 파일 {row_number}행에 이미 운송장번호가 있습니다. 중복 발송 여부를 먼저 확인하세요.")
        result.append(row)
    if not result:
        raise ValueError("주문 또는 운송장 데이터 행이 없습니다.")
    return result


def _ship_decrypt_naver(data):
    """네이버 암호화 파일을 먼저 읽는다. 1행 제거 여부는 복호화 후 헤더로 판단."""
    if msoffcrypto is None:
        raise RuntimeError("암호화 네이버 파일을 읽으려면 msoffcrypto-tool 라이브러리가 필요합니다.")
    try:
        crypt = msoffcrypto.OfficeFile(io.BytesIO(data))
        crypt.load_key(password=str(st.secrets.get("SMARTSTORE_PASSWORD", SMARTSTORE_PASSWORD)))
        output = io.BytesIO()
        crypt.decrypt(output)
        return output.getvalue()
    except Exception as exc:
        raise ValueError("네이버 엑셀 암호를 풀지 못했습니다. 비밀번호 1111 또는 파일을 확인해 주세요.") from exc


def _ship_read_file(filename, data):
    """4개 파일을 헤더로 판별. 암호화 네이버의 안내문 유무도 처리."""
    import zipfile
    from openpyxl import load_workbook
    if not filename.lower().endswith(".xlsx"):
        raise ValueError(f"{filename}: .xlsx 파일만 지원합니다.")
    if not data or len(data) > 25 * 1024 * 1024:
        raise ValueError(f"{filename}: 파일이 비었거나 25MB를 초과했습니다.")
    if data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        # 암호화 파일이 이미 1행 제거된 경우도 있다. 복호화 후 헤더를 검사한다.
        data = _ship_decrypt_naver(data)
    elif not data.startswith(b"PK"):
        raise ValueError(f"{filename}: 올바른 엑셀(.xlsx) 파일이 아닙니다.")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if sum(x.file_size for x in z.infolist()) > 180 * 1024 * 1024:
            raise ValueError(f"{filename}: 압축 해제 크기가 너무 큽니다.")
    wb = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
    try:
        if len(wb.worksheets) != 1:
            raise ValueError(f"{filename}: 시트가 여러 개라 자동 판별하지 않았습니다.")
        sheet = wb.worksheets[0]
        head1 = tuple(x.value for x in next(sheet.iter_rows(min_row=1, max_row=1)))
        columns = {str(x or "").strip() for x in head1}
        if {"상품주문번호", "송장번호", "수취인명"} <= columns:
            role = "naver_order"
        elif {"묶음배송번호", "운송장번호", "수취인이름"} <= columns:
            role = "coupang_order"
        elif {"운송장번호", "받는분성명", "받는분전화번호"} <= columns:
            role = "cj"
        elif len(head1) > 1:
            # 암호를 이미 해제했지만 아직 첫 행이 남아 있는 네이버 파일
            second = tuple(x.value for x in next(sheet.iter_rows(min_row=2, max_row=2)))
            if {"상품주문번호", "송장번호", "수취인명"} <= {str(x or "").strip() for x in second}:
                wb.close()
                data, _, _ = convert_naver_cj_excel(data)
                return _ship_read_file(filename, data)
            raise ValueError(f"{filename}: 네이버·쿠팡·CJ 중 지원되는 양식이 아닙니다.")
        else:
            raise ValueError(f"{filename}: 지원되는 양식이 아닙니다.")
        records = _ship_extract_records(sheet, head1, role)
        return {"file": filename, "role": role, "data": data, "records": records,
                "columns": len(head1), "tracking_column": list(head1).index("송장번호" if role == "naver_order" else "운송장번호") + 1}
    finally:
        wb.close()


def _ship_assign_carrier_files(order_files, carrier_files):
    """같은 형식의 CJ 파일 2개를 수취인 키의 겹침으로 플랫폼별로 구별."""
    if len(carrier_files) != 2:
        raise ValueError("CJ 송장번호 파일 2개가 필요합니다.")
    kinds = ("naver_order", "coupang_order")
    identity_sets = {role: {_ship_identity(r) for r in order_files[role]["records"]} for role in kinds}
    scores = []
    for cf in carrier_files:
        scores.append({role: sum(_ship_identity(r) in identity_sets[role] for r in cf["records"]) for role in kinds})
    forward = scores[0]["naver_order"] + scores[1]["coupang_order"]
    reverse = scores[0]["coupang_order"] + scores[1]["naver_order"]
    if forward == reverse or max(forward, reverse) == 0:
        raise ValueError("CJ 파일의 네이버·쿠팡 구분이 불확실합니다. 올바른 송장번호 파일 2개인지 확인하세요.")
    if forward > reverse:
        selected = {"naver_order": carrier_files[0], "coupang_order": carrier_files[1]}
    else:
        selected = {"naver_order": carrier_files[1], "coupang_order": carrier_files[0]}
    return selected


def _ship_match_platform(order_file, cj_file):
    """이름+전화+주소 정확히 일치. 복수 송장은 옵션+수량으로 추가 구분."""
    from collections import defaultdict, Counter
    order_groups, cj_groups = defaultdict(list), defaultdict(list)
    for r in order_file["records"]:
        order_groups[_ship_identity(r)].append(r)
    for r in cj_file["records"]:
        cj_groups[_ship_identity(r)].append(r)
    issues, assignments = [], {}
    track_to_recipient = {}
    for identity, group in cj_groups.items():
        for row in group:
            tracking = row["tracking_digits"]
            if tracking in track_to_recipient and track_to_recipient[tracking] != identity:
                issues.append(f"CJ {row['row']}행: 다른 수취인에게 동일 송장번호가 사용됐습니다.")
            track_to_recipient[tracking] = identity
        if identity not in order_groups:
            issues.append(f"CJ {group[0]['row']}행: 주문 엑셀에서 수취인 정보가 일치하지 않습니다.")

    for identity, order_group in order_groups.items():
        cj_group = cj_groups.get(identity, [])
        if not cj_group:
            issues.append(f"주문 {order_group[0]['row']}행: CJ 송장에 일치하는 수취인이 없습니다.")
            continue
        distinct_tracks = {r["tracking_digits"] for r in cj_group}
        if len(order_group) != len(cj_group):
            issues.append(f"주문 {order_group[0]['row']}행: 상품행 수가 CJ와 다릅니다 ({len(order_group)}행/{len(cj_group)}행).")
            continue
        order_sig = Counter(_ship_signature(x) for x in order_group)
        cj_sig = Counter(_ship_signature(x) for x in cj_group)
        if order_sig != cj_sig:
            issues.append(f"주문 {order_group[0]['row']}행: CJ의 상품 옵션·수량과 일치하지 않습니다.")
            continue
        if len(distinct_tracks) == 1:
            for r in order_group:
                assignments[r["row"]] = cj_group[0]["tracking"]
        else:
            # 같은 수취인에 송장이 여러 개인 경우 옵션+수량으로만 연결
            by_sig = defaultdict(list)
            for r in cj_group:
                by_sig[_ship_signature(r)].append(r)
            for order in order_group:
                options = by_sig[_ship_signature(order)]
                if len(options) != 1:
                    issues.append(f"주문 {order['row']}행: 동일 옵션에 여러 송장번호가 있어 자동 매칭할 수 없습니다.")
                    continue
                assignments[order["row"]] = options[0]["tracking"]
    return assignments, issues


def _ship_build_completed(order_file, assignments):
    """기존 셀은 그대로 유지하고 송장번호 열만 변경한 뒤 재검증."""
    from openpyxl import load_workbook
    baseline = load_workbook(io.BytesIO(order_file["data"]), read_only=True, data_only=False)
    try:
        original_values = [tuple(row) for row in baseline.worksheets[0].values]
    finally:
        baseline.close()
    if len(assignments) != len(order_file["records"]):
        raise ValueError("송장 매칭되지 않은 주문행이 있어 완성 파일 생성을 중단했습니다.")
    wb = load_workbook(io.BytesIO(order_file["data"]), keep_links=True)
    try:
        ws = wb.worksheets[0]
        col = order_file["tracking_column"]
        for row_number, tracking in assignments.items():
            ws.cell(row_number, col).value = tracking  # 문자형으로 저장 (앞자리 0 보존)
            ws.cell(row_number, col).number_format = "@"
        output = io.BytesIO()
        wb.save(output)
        result = output.getvalue()
    finally:
        wb.close()
    check = load_workbook(io.BytesIO(result), read_only=True, data_only=False)
    try:
        actual_values = [tuple(row) for row in check.worksheets[0].values]
    finally:
        check.close()
    if len(original_values) != len(actual_values):
        raise RuntimeError("완성 엑셀의 행 수가 원본과 달라졌습니다.")
    for row_number, (original_row, completed_row) in enumerate(zip(original_values, actual_values), start=1):
        if len(original_row) != len(completed_row):
            raise RuntimeError(f"{row_number}행의 열 수가 변경됐습니다.")
        for col_number, (before, after) in enumerate(zip(original_row, completed_row), start=1):
            if row_number in assignments and col_number == order_file["tracking_column"]:
                if _ship_digits(after) != _ship_digits(assignments[row_number]):
                    raise RuntimeError(f"{row_number}행 송장번호 검증에 실패했습니다.")
            elif before != after and not (before in (None, "") and after in (None, "")):
                raise RuntimeError(f"원본의 송장번호 이외 값이 변경됐습니다 ({row_number}행).")
    return result


def _ship_naver_dispatch_xls(order_file, assignments):
    """네이버 발송처리와 CRM 분석을 동시에 지원하는 비암호화 BIFF8 .xls.

    필수 4개 열과 `발송처리` 시트 이름을 유지하되 원본 72개 열을 보존한다.
    따라서 최종 송장 파일에서도 주문일/고객/상품/옵션/수량을 다시 분석할 수 있다.
    """
    try:
        import xlwt
    except ImportError as exc:
        raise RuntimeError(
            "네이버 .xls 생성 라이브러리(xlwt)가 없습니다. "
            "GitHub requirements.txt에 xlwt==1.3.0을 추가해 주세요."
        ) from exc
    from openpyxl import load_workbook

    source = load_workbook(io.BytesIO(order_file["data"]), read_only=True, data_only=True)
    try:
        ws = source.worksheets[0]
        heading = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        fields = ("상품주문번호", "배송방법", "택배사", "송장번호")
        missing = [field for field in fields if field not in heading]
        if missing:
            raise ValueError("네이버 주문 엑셀 필수 열 누락: " + ", ".join(missing))
        if len(heading) > 256:
            raise ValueError("XLS 최대 256열을 초과해 주문 원본을 보존할 수 없습니다.")
        idx = {field: heading.index(field) for field in fields}
        expected_rows = {row["row"] for row in order_file["records"]}
        if expected_rows != set(assignments):
            raise ValueError("송장번호가 모든 네이버 상품주문행에 매칭되지 않았습니다.")
        if len(expected_rows) >= 65536:
            raise ValueError("XLS 형식은 65,535개 주문행까지만 보존할 수 있습니다.")

        result_book = xlwt.Workbook(encoding="utf-8")
        target = result_book.add_sheet("발송처리")
        date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
        for col, title in enumerate(heading):
            target.write(0, col, title)
            if col in idx.values():
                target.col(col).width = min(34, max(18, len(title) * 3)) * 256

        written = 0
        seen_ids = set()
        for source_row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if source_row_number not in expected_rows:
                continue
            cells = list(values)
            order_id = cells[idx["상품주문번호"]]
            if isinstance(order_id, float):
                if not order_id.is_integer():
                    raise ValueError(f"네이버 {source_row_number}행 상품주문번호가 유효하지 않습니다.")
                order_id = int(order_id)
            order_id = str(order_id or "").strip()
            if not order_id.isdigit() or not 10 <= len(order_id) <= 22:
                raise ValueError(f"네이버 {source_row_number}행 상품주문번호를 확인해 주세요.")
            if order_id in seen_ids:
                raise ValueError(f"상품주문번호 {order_id} 중복으로 생성을 중단했습니다.")
            seen_ids.add(order_id)
            tracking = _ship_digits(assignments[source_row_number])
            if not 8 <= len(tracking) <= 20:
                raise ValueError(f"네이버 {source_row_number}행 송장번호를 확인해 주세요.")
            # 원본 전체 주문 정보는 보존. 네이버 발송처리 필수 4개 열만 변경.
            cells[idx["상품주문번호"]] = order_id
            cells[idx["배송방법"]] = "택배,등기,소포"
            cells[idx["택배사"]] = "CJ대한통운"
            cells[idx["송장번호"]] = tracking
            for col, value in enumerate(cells[:len(heading)]):
                if value is None:
                    continue
                # 전화번호·주문번호처럼 긴 숫자는 엑셀 15자리 숫자 정밀도 손실을 막는다.
                title = heading[col]
                if col in (idx["상품주문번호"], idx["송장번호"]) or any(
                    marker in title for marker in ("주문번호", "연락처", "전화번호", "우편번호", "송장번호", "상품번호", "묶음배송번호")
                ):
                    target.write(written + 1, col, str(value))
                elif isinstance(value, (datetime, date)):
                    target.write(written + 1, col, value, date_style)
                elif isinstance(value, (str, int, float, bool)):
                    target.write(written + 1, col, value)
                else:
                    target.write(written + 1, col, str(value))
            written += 1
        if written != len(expected_rows):
            raise RuntimeError("네이버 발송처리 파일의 주문행 수가 원본과 다릅니다.")
        output = io.BytesIO()
        result_book.save(output)
        payload = output.getvalue()
        if not payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise RuntimeError("XLS 파일 생성에 실패했습니다.")
        return payload
    finally:
        source.close()

def make_shipment_zip(files, yyyymmdd):
    """파일 순서 무관. 네이버는 주문원본까지 담은 일괄발송용 xls, 쿠팡은 xlsx를 ZIP으로 반환."""
    import zipfile
    if len(files) != 4:
        raise ValueError(f"파일 4개가 필요합니다. 현재 {len(files)}개입니다.")
    if sum(len(b) for _, b in files) > 70 * 1024 * 1024:
        raise ValueError("전체 업로드 용량이 70MB를 초과합니다.")
    parsed = [_ship_read_file(name, data) for name, data in files]
    order_files = {}
    cj_files = []
    for f in parsed:
        if f["role"] == "cj":
            cj_files.append(f)
        elif f["role"] in order_files:
            raise ValueError(f"같은 플랫폼의 주문 엑셀이 중복됐습니다: {f['role']}")
        else:
            order_files[f["role"]] = f
    if set(order_files) != {"naver_order", "coupang_order"} or len(cj_files) != 2:
        raise ValueError("네이버 주문 1개, 쿠팡 주문 1개, CJ 송장 2개를 넣어주세요.")
    cj_map = _ship_assign_carrier_files(order_files, cj_files)
    built, summary, issues = {}, {}, []
    for platform, label in [("naver_order", "네이버"), ("coupang_order", "쿠팡")]:
        order, cj = order_files[platform], cj_map[platform]
        assignments, errs = _ship_match_platform(order, cj)
        summary[label] = {"orders": len(order["records"]), "carriers": len(cj["records"]),
                          "matched": len(assignments), "unique_tracks": len({x["tracking_digits"] for x in cj["records"]}),
                          "order_file": order["file"], "cj_file": cj["file"]}
        if errs:
            issues.extend(f"{label}: {err}" for err in errs)
        if len(assignments) != len(order["records"]):
            issues.append(f"{label}: 매칭되지 않은 주문행이 있습니다.")
        if not errs and len(assignments) == len(order["records"]):
            built[label] = (_ship_naver_dispatch_xls(order, assignments) if label == "네이버"
                            else _ship_build_completed(order, assignments))
    if issues:
        return None, summary, issues
    output = io.BytesIO()
    names = {"네이버": f"★네이버 완성({yyyymmdd}).xls", "쿠팡": f"★쿠팡 완성({yyyymmdd}).xlsx"}
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for platform in ("네이버", "쿠팡"):
            archive.writestr(names[platform], built[platform])
    return output.getvalue(), summary, []


def detect_market(filename, df):
    name = filename.lower()
    cols = set(map(str, df.columns))

    if "delivery" in name or "coupang" in name or "쿠팡" in name:
        return "쿠팡"
    if "스마트스토어" in name or "smartstore" in name or "naver" in name or "네이버" in name:
        return "네이버"
    if "묶음배송번호" in cols or "노출상품명(옵션명)" in cols:
        return "쿠팡"
    if "상품주문번호" in cols or "통합배송지" in cols or "수취인연락처1" in cols:
        return "네이버"

    return "기타"


def standardize(df, market, filename):
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]

    if market == "네이버":
        order_unit_col = pick_col(df, ["주문번호"])
        product_order_col = pick_col(df, ["상품주문번호"])
        order_date_col = pick_col(df, ["주문일시", "결제일", "발주확인일", "송장출력일", "발송일"])
        amount_col = pick_col(df, ["최종 상품별 총 주문금액", "최초 상품별 총 주문금액", "정산예정금액"])
        address_col = pick_col(df, ["통합배송지", "기본배송지", "배송지주소", "수취인주소"])
        product_col = pick_col(df, ["상품명"])
    elif market == "쿠팡":
        order_unit_col = pick_col(df, ["묶음배송번호", "주문번호"])
        product_order_col = pick_col(df, ["주문번호"])
        order_date_col = pick_col(df, ["주문일", "출고일(발송일)", "배송완료일"])
        amount_col = pick_col(df, ["결제액", "옵션판매가(판매단가)"])
        address_col = pick_col(df, ["수취인 주소", "배송지주소", "수취인주소", "주소"])
        product_col = pick_col(df, ["노출상품명(옵션명)", "등록상품명"])
    else:
        order_unit_col = pick_col(df, ["주문번호", "묶음배송번호", "상품주문번호"])
        product_order_col = pick_col(df, ["상품주문번호", "주문번호"])
        order_date_col = pick_col(df, ["주문일시", "주문일", "결제일"])
        amount_col = pick_col(df, ["결제금액", "결제액", "최종 상품별 총 주문금액"])
        address_col = pick_col(df, ["통합배송지", "수취인 주소", "주소"])
        product_col = pick_col(df, ["상품명", "노출상품명(옵션명)", "등록상품명"])

    receiver_col = pick_col(df, ["수취인명", "수취인이름", "수령인", "받는분", "수령자명"])
    receiver_phone_col = pick_col(df, ["수취인연락처1", "수취인전화번호", "수취인연락처", "수령인전화번호"])
    buyer_col = pick_col(df, ["구매자명", "구매자", "주문자명"])
    buyer_phone_col = pick_col(df, ["구매자연락처", "구매자전화번호", "주문자연락처"])
    qty_col = pick_col(df, ["수량", "구매수(수량)", "구매수량"])
    option_col = pick_col(df, ["옵션정보", "옵션명", "옵션값", "옵션내용"])

    out = pd.DataFrame(index=df.index)
    out["channel"] = market
    out["source_file"] = filename

    out["order_no"] = safe_series(df, order_unit_col, "").astype(str).str.replace(".0", "", regex=False)
    out["product_order_no"] = safe_series(df, product_order_col, "").astype(str).str.replace(".0", "", regex=False)
    out["order_no"] = out["order_no"].where(out["order_no"].str.strip() != "", out["product_order_no"])

    out["order_date"] = pd.to_datetime(safe_series(df, order_date_col, pd.NaT), errors="coerce")
    out["receiver_name"] = safe_series(df, receiver_col, "").map(normalize_text)
    out["receiver_phone"] = safe_series(df, receiver_phone_col, "").map(normalize_phone)
    out["buyer_name"] = safe_series(df, buyer_col, "").map(normalize_text)
    out["buyer_phone"] = safe_series(df, buyer_phone_col, "").map(normalize_phone)
    out["address"] = safe_series(df, address_col, "").map(normalize_address)
    product_base = safe_series(df, product_col, "").map(normalize_text)
    option_text = safe_series(df, option_col, "").map(normalize_text)
    if market == "네이버":
        out["product_name"] = option_text.where(option_text.str.strip() != "", product_base)
    else:
        out["product_name"] = (product_base + " " + option_text).str.strip()
    out["quantity"] = pd.to_numeric(safe_series(df, qty_col, 1), errors="coerce").fillna(1).astype(int)
    out["amount"] = pd.to_numeric(safe_series(df, amount_col, 0), errors="coerce").fillna(0).astype(int)

    name_key = out["receiver_name"].where(out["receiver_name"] != "", out["buyer_name"])
    phone_key = out["receiver_phone"].where(out["receiver_phone"] != "", out["buyer_phone"])

    out["customer_key"] = name_key.fillna("").astype(str).str.strip() + "|" + out["address"].fillna("").astype(str).str.strip()
    out["customer_key_strict"] = (
        name_key.fillna("").astype(str).str.strip()
        + "|" + phone_key.fillna("").astype(str).str.strip()
        + "|" + out["address"].fillna("").astype(str).str.strip()
    )

    out["order_key"] = out["channel"] + "|" + out["order_no"].astype(str)
    out = out[(name_key != "") | (phone_key != "") | (out["address"] != "")]
    out = out.dropna(subset=["order_date"])
    return out


def join_unique(series):
    vals = []
    for v in series:
        v = str(v).strip()
        if v and v.lower() != "nan" and v not in vals:
            vals.append(v)
    return ", ".join(vals)


def grade(n):
    if n >= 5:
        return "VIP"
    if n >= 3:
        return "우수고객"
    if n >= 2:
        return "재구매"
    return "신규"


def churn(days):
    if days >= 180:
        return "장기이탈"
    if days >= 90:
        return "이탈위험"
    if days >= 60:
        return "관심필요"
    return "정상"


def make_order_level(raw_orders, use_strict=False):
    key_col = "customer_key_strict" if use_strict else "customer_key"

    grouped = (
        raw_orders.groupby(["channel", "order_no", "order_key", key_col, "receiver_name", "receiver_phone", "buyer_name", "buyer_phone", "address"], dropna=False)
        .agg(
            order_date=("order_date", "min"),
            product_order_nos=("product_order_no", join_unique),
            product_names=("product_name", join_unique),
            total_quantity=("quantity", "sum"),
            total_amount=("amount", "sum"),
            source_files=("source_file", join_unique),
            raw_rows=("product_order_no", "count"),
        )
        .reset_index()
        .rename(columns={key_col: "analysis_customer_key"})
    )
    return grouped


@st.cache_data(ttl=180, max_entries=4, show_spinner=False)
def make_customer_df(order_level):
    if order_level.empty:
        return pd.DataFrame()

    customer = (
        order_level.groupby("analysis_customer_key", dropna=False)
        .agg(
            customer_name=("receiver_name", "first"),
            phone=("receiver_phone", "first"),
            address=("address", "first"),
            order_count=("order_key", pd.Series.nunique),
            first_order_date=("order_date", "min"),
            last_order_date=("order_date", "max"),
            total_amount=("total_amount", "sum"),
            total_quantity=("total_quantity", "sum"),
            channels=("channel", join_unique),
            products=("product_names", join_unique),
        )
        .reset_index()
    )

    # Supabase timestamptz는 timezone-aware 값으로 들어오므로 UTC 기준으로 통일
    today = pd.Timestamp.now(tz="UTC").normalize()
    customer["grade"] = customer["order_count"].apply(grade)
    customer["customer_type"] = customer["order_count"].apply(lambda x: "재구매" if x >= 2 else "신규")
    last_order = pd.to_datetime(customer["last_order_date"], utc=True, errors="coerce").dt.normalize()
    customer["days_since_last_order"] = (today - last_order).dt.days.fillna(0).astype(int)
    customer["churn_status"] = customer["days_since_last_order"].apply(churn)
    customer["avg_order_amount"] = (customer["total_amount"] / customer["order_count"]).round(0).fillna(0).astype(int)
    return customer


@st.cache_data(ttl=120, show_spinner=False)
def fetch_all(table_name):
    if sb is None:
        return pd.DataFrame()

    rows = []
    step = 1000
    start = 0

    # ORDER BY 없는 페이지 조회는 1만 건 이상에서 중복/누락 위험이 있다.
    sort_columns = {"orders": "order_key", "blacklist": "customer_key"}
    sort_column = sort_columns.get(table_name)
    while True:
        query = sb.table(table_name).select("*")
        if sort_column:
            query = query.order(sort_column)
        # 네트워크의 일시적인 DNS/연결 오류는 짧게 재시도하고,
        # 계속 실패하면 개인정보/API 오류문 대신 원인 종류만 보여준다.
        res = None
        for attempt in range(3):
            try:
                res = query.range(start, start + step - 1).execute()
                break
            except Exception as exc:
                if attempt == 2:
                    raise RuntimeError(
                        f"Supabase {table_name} 목록 조회 실패 ({type(exc).__name__}). "
                        "프로젝트 상태와 네트워크 연결을 확인해 주세요."
                    ) from None
                time.sleep((attempt + 1) * 0.35)
        data = res.data or []
        rows.extend(data)
        if len(data) < step:
            break
        start += step

    return pd.DataFrame(rows)


def upsert_orders(order_level):
    if sb is None or order_level.empty:
        return 0

    records = []
    now = datetime.now().isoformat(timespec="seconds")

    for _, r in order_level.iterrows():
        records.append({
            "order_key": str(r["order_key"]),
            "channel": str(r["channel"]),
            "order_no": str(r["order_no"]),
            "customer_key": str(r["analysis_customer_key"]),
            "receiver_name": str(r["receiver_name"]),
            "receiver_phone": str(r["receiver_phone"]),
            "buyer_name": str(r["buyer_name"]),
            "buyer_phone": str(r["buyer_phone"]),
            "address": str(r["address"]),
            "order_date": pd.to_datetime(r["order_date"], utc=True, errors="coerce").isoformat(),
            "product_order_nos": str(r["product_order_nos"]),
            "product_names": str(r["product_names"]),
            "total_quantity": int(r["total_quantity"]),
            "total_amount": int(r["total_amount"]),
            "source_files": str(r["source_files"]),
            "raw_rows": int(r["raw_rows"]),
            "updated_at": now,
        })

    # 같은 주문번호가 여러 파일에 중복 업로드돼도 DB 저장 전에 1개만 남김
    dedup = {}
    for rec in records:
        dedup[rec["order_key"]] = rec
    records = list(dedup.values())

    def save_chunk_safely(chunk):
        """큰 묶음 저장 실패 시 자동으로 반씩 쪼개 저장"""
        if not chunk:
            return 0

        try:
            sb.table("orders").upsert(chunk, on_conflict="order_key").execute()
            return len(chunk)
        except Exception as e:
            # 연결/DNS/권한 장애는 데이터를 잘게 쪼개도 해결되지 않는다.
            # 무한에 가까운 재시도(1000건→1건)를 피하고 즉시 안전하게 중단한다.
            error_type = type(e).__name__.lower()
            description = str(e).lower()[:500]
            no_split = (
                any(term in error_type for term in ("connect", "timeout", "network"))
                or any(term in description for term in (
                    "name or service not known", "temporary failure", "connection refused",
                    "connection error", "timed out", "timeout", "network error", "dns",
                    "401 unauthorized", "403 forbidden", "rate limit", "too many requests",
                    "service unavailable", "bad gateway", "gateway timeout"
                ))
            )
            if no_split:
                raise RuntimeError(
                    f"Supabase 주문 저장 연결 오류 ({type(e).__name__}). "
                    "일부 앞선 주문은 저장됐을 수 있습니다. 상태를 확인한 후 다시 저장하세요."
                ) from None
            if len(chunk) <= 1:
                # 어떤 주문 1건 자체가 문제일 때만 화면에 명확히 표시
                raise RuntimeError(
                    f"주문 1건 저장 실패 ({type(e).__name__}). 앞서 성공한 묶음은 저장됐을 수 있습니다. "
                    "실패 원인을 확인한 뒤 중복 업로드 대신 해당 주문을 점검해 주세요."
                ) from None

            mid = len(chunk) // 2
            return save_chunk_safely(chunk[:mid]) + save_chunk_safely(chunk[mid:])

    # 기본은 1000건씩 시도, 실패하면 자동 분할
    chunk_size = 1000
    saved_count = 0

    for i in range(0, len(records), chunk_size):
        chunk = records[i:i + chunk_size]
        saved_count += save_chunk_safely(chunk)

    fetch_all.clear()
    st.session_state.pop("_crm_excel_bytes_v21", None)
    st.session_state["_orders_revision_v21"] = st.session_state.get("_orders_revision_v21", 0) + 1
    return saved_count


def fetch_orders_as_order_level():
    df = fetch_all("orders")
    if df.empty:
        return df

    df["order_date"] = pd.to_datetime(df["order_date"], errors="coerce")
    df = df.rename(columns={
        "customer_key": "analysis_customer_key",
        "receiver_name": "receiver_name",
        "receiver_phone": "receiver_phone",
        "buyer_name": "buyer_name",
        "buyer_phone": "buyer_phone",
        "address": "address",
        "product_order_nos": "product_order_nos",
        "product_names": "product_names",
        "total_quantity": "total_quantity",
        "total_amount": "total_amount",
        "source_files": "source_files",
        "raw_rows": "raw_rows",
    })
    return df


def fetch_blacklist():
    df = fetch_all("blacklist")
    if df.empty:
        return pd.DataFrame(columns=["customer_key", "customer_name", "phone", "address", "status", "reason", "memo", "incident_date", "result", "updated_at"])
    return df


def save_blacklist(row):
    if sb is None:
        return

    row["updated_at"] = datetime.now().isoformat(timespec="seconds")
    sb.table("blacklist").upsert([row], on_conflict="customer_key").execute()
    fetch_all.clear()
    st.session_state.pop("_crm_excel_bytes_v21", None)
    st.session_state["_crm_customer_flags_revision_v21"] = st.session_state.get("_crm_customer_flags_revision_v21", 0) + 1


def delete_blacklist(customer_key):
    if sb is None:
        return
    sb.table("blacklist").delete().eq("customer_key", customer_key).execute()
    fetch_all.clear()
    st.session_state.pop("_crm_excel_bytes_v21", None)
    st.session_state["_crm_customer_flags_revision_v21"] = st.session_state.get("_crm_customer_flags_revision_v21", 0) + 1




def fetch_customer_notes():
    df = fetch_blacklist()
    if df.empty:
        return pd.DataFrame(columns=["customer_key", "memo", "updated_at"])
    notes = df[df.get("reason", "") == "고객메모"].copy() if "reason" in df.columns else pd.DataFrame()
    if notes.empty:
        return pd.DataFrame(columns=["customer_key", "memo", "updated_at"])
    return notes[["customer_key", "memo", "updated_at"]].copy()


def save_customer_note(customer_key, customer_name, phone, address, memo):
    if sb is None:
        return
    row = {
        "customer_key": str(customer_key),
        "customer_name": str(customer_name),
        "phone": str(phone),
        "address": str(address),
        "status": "메모",
        "reason": "고객메모",
        "memo": str(memo),
        "incident_date": str(date.today()),
        "result": "",
        "updated_at": datetime.now().isoformat(timespec="seconds"),
    }
    sb.table("blacklist").upsert([row], on_conflict="customer_key").execute()
    fetch_all.clear()
    st.session_state.pop("_crm_excel_bytes_v21", None)
    st.session_state["_crm_customer_flags_revision_v21"] = st.session_state.get("_crm_customer_flags_revision_v21", 0) + 1


@st.cache_data(ttl=300, max_entries=3, show_spinner=False)
def _excel_safe_value(value):
    """사용자 입력 문자열이 Excel에서 계산식으로 실행되지 않게 텍스트로 내보낸다."""
    if isinstance(value, str):
        # leading whitespace / control characters can be ignored by spreadsheet viewers.
        if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def to_excel_bytes(sheets):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, df in sheets.items():
            safe_df = df.copy()
            # Excel은 timezone-aware datetime을 저장하지 못하므로 timezone 제거
            for col in safe_df.columns:
                if pd.api.types.is_datetime64_any_dtype(safe_df[col]):
                    try:
                        safe_df[col] = safe_df[col].dt.tz_localize(None)
                    except Exception:
                        safe_df[col] = pd.to_datetime(safe_df[col], errors="coerce", utc=True).dt.tz_localize(None)
                elif safe_df[col].dtype == "object":
                    # 문자열 안에 섞인 UTC datetime도 안전하게 문자열 처리
                    safe_df[col] = safe_df[col].apply(lambda x: x.isoformat() if hasattr(x, "isoformat") else x)
            # 고객 이름/주소/메모/상품명 등도 스프레드시트 수식으로 해석되지 않도록 보호.
            for col in safe_df.columns:
                if safe_df[col].dtype == "object" or pd.api.types.is_string_dtype(safe_df[col]):
                    safe_df[col] = safe_df[col].map(_excel_safe_value)
            safe_df.to_excel(writer, index=False, sheet_name=name[:31])
    output.seek(0)
    return output.getvalue()




def analyze_today_orders(order_preview, orders_db, customer_df, black_df):
    """업로드한 주문을 DB 저장 전 상태에서 기존 DB와 비교"""
    if order_preview is None or order_preview.empty:
        return pd.DataFrame(), pd.DataFrame()

    today_orders = order_preview.copy()
    today_orders["order_date"] = pd.to_datetime(today_orders["order_date"], utc=True, errors="coerce")

    if customer_df is None or customer_df.empty:
        existing = pd.DataFrame(columns=[
            "analysis_customer_key", "order_count", "first_order_date", "last_order_date",
            "grade", "days_since_last_order", "total_amount", "avg_order_amount"
        ])
    else:
        need_cols = [
            "analysis_customer_key", "order_count", "first_order_date", "last_order_date",
            "grade", "days_since_last_order", "total_amount", "avg_order_amount"
        ]
        existing = customer_df.copy()
        for c in need_cols:
            if c not in existing.columns:
                existing[c] = None
        existing = existing[need_cols]

    today_customer = (
        today_orders.groupby("analysis_customer_key", dropna=False)
        .agg(
            오늘고객명=("receiver_name", "first"),
            오늘전화번호=("receiver_phone", "first"),
            오늘주소=("address", "first"),
            오늘주문수=("order_key", pd.Series.nunique),
            오늘상품=("product_names", join_unique),
            오늘수량=("total_quantity", "sum"),
            오늘금액=("total_amount", "sum"),
            오늘채널=("channel", join_unique),
            오늘첫주문일=("order_date", "min"),
            오늘주문번호=("order_no", join_unique),
        )
        .reset_index()
    )

    result = today_customer.merge(existing, on="analysis_customer_key", how="left")
    result["기존주문횟수"] = result["order_count"].fillna(0).astype(int)
    result["이번포함총주문"] = result["기존주문횟수"] + result["오늘주문수"].fillna(0).astype(int)
    result["첫주문일"] = result["first_order_date"]
    result["최근주문일"] = result["last_order_date"]

    def status_row(r):
        old_count = int(r["기존주문횟수"])
        old_grade = str(r.get("grade", ""))
        days = r.get("days_since_last_order", None)

        if old_count <= 0:
            return "⚪ 신규"
        if old_grade == "VIP" or old_count >= 5:
            return "🟡 VIP 재주문"
        try:
            if pd.notna(days) and int(days) >= 90:
                return "🔵 복귀 재주문"
        except Exception:
            pass
        return "🟢 재주문"

    result["오늘고객상태"] = result.apply(status_row, axis=1)

    if black_df is not None and not black_df.empty:
        b = black_df.copy()
        for c in ["customer_key", "status", "reason", "memo", "incident_date", "result"]:
            if c not in b.columns:
                b[c] = ""
        b = b[["customer_key", "status", "reason", "memo", "incident_date", "result"]]
        result = result.merge(b, left_on="analysis_customer_key", right_on="customer_key", how="left")
    else:
        result["status"] = ""
        result["reason"] = ""
        result["memo"] = ""
        result["incident_date"] = ""
        result["result"] = ""

    result["블랙상태"] = result["status"].fillna("정상").replace("", "정상")
    result.loc[result["블랙상태"].isin(["주의", "블랙"]), "오늘고객상태"] = "🔴 주의/블랙"

    final_cols = [
        "오늘고객상태", "오늘고객명", "오늘전화번호", "오늘채널", "오늘상품",
        "오늘주문수", "기존주문횟수", "이번포함총주문",
        "첫주문일", "최근주문일", "days_since_last_order",
        "오늘금액", "오늘수량", "블랙상태", "reason", "memo", "result",
        "오늘주소", "오늘주문번호", "analysis_customer_key"
    ]

    for c in final_cols:
        if c not in result.columns:
            result[c] = ""

    result = result[final_cols].rename(columns={
        "days_since_last_order": "최근구매후경과일",
        "reason": "블랙사유",
        "memo": "블랙메모",
        "result": "처리결과",
        "analysis_customer_key": "고객키",
    })

    result = result.sort_values(
        by=["오늘고객상태", "이번포함총주문", "오늘금액"],
        ascending=[True, False, False]
    )

    detail = today_orders.merge(
        result[["고객키", "오늘고객상태", "기존주문횟수", "이번포함총주문", "블랙상태", "블랙사유"]],
        left_on="analysis_customer_key",
        right_on="고객키",
        how="left"
    )

    return result, detail



def style_customer_crm_table(df):
    if df is None or df.empty:
        return df

    def row_style(row):
        black = str(row.get("black_status", row.get("status", "")))
        grade_val = str(row.get("grade", ""))
        order_count = int(row.get("order_count", 0)) if pd.notna(row.get("order_count", 0)) else 0

        if black == "블랙":
            return ["background-color: #000000; color: #ffffff; font-weight: 800"] * len(row)
        if black == "주의":
            return ["background-color: #fee2e2; color: #111827; font-weight: 600"] * len(row)
        if grade_val == "VIP" or order_count >= 5:
            return ["background-color: #fef3c7; color: #111827; font-weight: 700"] * len(row)
        if order_count >= 2:
            return ["background-color: #ffedd5; color: #111827; font-weight: 600"] * len(row)
        return ["background-color: #dcfce7; color: #111827; font-weight: 600"] * len(row)

    return df.style.apply(row_style, axis=1)

def style_today_customer_table(df):
    """오늘 주문 분석 표 색상"""
    if df is None or df.empty:
        return df

    def row_style(row):
        status = str(row.get("오늘고객상태", ""))
        black = str(row.get("블랙상태", ""))

        if "🔴" in status or black in ["주의", "블랙"]:
            return ["background-color: #fee2e2; color: #111827; font-weight: 600"] * len(row)
        if "🟡" in status:
            return ["background-color: #fef3c7; color: #111827; font-weight: 700"] * len(row)
        if "🔵" in status:
            return ["background-color: #dbeafe"] * len(row)
        if "🟢" in status:
            return ["background-color: #dcfce7; color: #111827; font-weight: 600"] * len(row)
        return [""] * len(row)

    return df.style.apply(row_style, axis=1)


def make_ai_context(customer_df, orders_db):
    """AI에게 넘길 CRM 요약. 개인정보 보호를 위해 이름/전화번호/주소는 전송하지 않음."""
    if customer_df is None or customer_df.empty:
        return "저장된 고객 데이터가 없습니다."

    total_orders = len(orders_db) if orders_db is not None else 0
    total_customers = len(customer_df)

    repeat_customers = int((customer_df["order_count"] >= 2).sum()) if "order_count" in customer_df.columns else 0
    vip_customers = int((customer_df["order_count"] >= 5).sum()) if "order_count" in customer_df.columns else 0
    repeat_rate = round(repeat_customers / total_customers * 100, 1) if total_customers else 0

    one_time = int((customer_df["order_count"] == 1).sum()) if "order_count" in customer_df.columns else 0
    two_time = int((customer_df["order_count"] == 2).sum()) if "order_count" in customer_df.columns else 0
    three_plus = int((customer_df["order_count"] >= 3).sum()) if "order_count" in customer_df.columns else 0

    if "order_count" in customer_df.columns:
        top_order_counts = customer_df["order_count"].value_counts().sort_index().head(20).to_dict()
    else:
        top_order_counts = {}

    product_text = ""
    if orders_db is not None and not orders_db.empty and "product_names" in orders_db.columns:
        # 원본 상품명/옵션 문자열에는 판매자가 입력한 개인 정보가 포함될 수 있다.
        # AI에는 식혜 5개 대분류의 주문 빈도만 전송하고 원본 텍스트는 전송하지 않는다.
        counts = orders_db["product_names"].map(classify_label).fillna("기타/미분류").value_counts()
        product_text = "\n".join(
            f"- {LABELS.get(item, '기타/미분류')}: {int(total):,}건"
            for item, total in counts.items()
        )

    monthly_text = ""
    if orders_db is not None and not orders_db.empty and "order_date" in orders_db.columns:
        temp = orders_db.copy()
        temp["월"] = pd.to_datetime(temp["order_date"], errors="coerce", utc=True).dt.to_period("M").astype(str)
        monthly = temp.groupby("월").size().tail(12)
        monthly_text = "\n".join([f"- {idx}: {val}건" for idx, val in monthly.items()])

    return f"""
식혜명가 CRM 집계 요약:
- 누적 주문수: {total_orders:,}건
- 누적 고객수: {total_customers:,}명
- 1회 구매 고객수: {one_time:,}명
- 2회 구매 고객수: {two_time:,}명
- 3회 이상 구매 고객수: {three_plus:,}명
- 재구매 고객수: {repeat_customers:,}명
- 재구매율: {repeat_rate}%
- VIP 고객수(5회 이상): {vip_customers:,}명

구매횟수 분포:
{top_order_counts}

최근 12개월 주문수:
{monthly_text}

주요 상품 빈도:
{product_text}

주의:
- 이 요약에는 고객명, 전화번호, 주소 등 개인정보를 포함하지 않는다.
- 답변도 개인정보 없이 집계와 전략 중심으로 한다.
"""


def auto_ai_crm_analysis(customer_df, orders_db):
    """버튼 한 번으로 CRM 자동 분석"""
    question = """
현재 CRM 데이터를 바탕으로 아래 항목을 실무적으로 분석해줘.

1. 현재 고객 구조 요약
2. 재구매율 상태 평가
3. VIP 고객 관리 포인트
4. 신규 고객을 재구매로 전환하기 위한 전략
5. 당장 실행할 마케팅 액션 TOP 5
6. 문자/쿠폰을 보낸다면 어떤 고객군을 우선해야 하는지
7. 사장님이 이번 주에 확인해야 할 체크리스트

답변은 너무 길지 않게, 바로 실행 가능한 형태로 정리해줘.
"""
    return ask_ai_crm(question, customer_df, orders_db)


def _ai_settings():
    """AI 설정 진단. 키는 UI/로그에 절대 표시하지 않음."""
    preferred = str(st.secrets.get("AI_PROVIDER", "")).strip().lower()
    gemini_key = str(st.secrets.get("GEMINI_API_KEY", "")).strip()
    openai_key = str(st.secrets.get("OPENAI_API_KEY", "")).strip()

    # 과거 버전은 AI_PROVIDER를 생략하면 OpenAI로 고정돼 Gemini 키가 무시됐음.
    if preferred not in ("gemini", "openai"):
        preferred = "gemini" if gemini_key else "openai"
    if preferred == "openai" and not openai_key and gemini_key:
        preferred = "gemini"
    if preferred == "gemini" and not gemini_key and openai_key:
        preferred = "openai"

    if preferred == "gemini":
        configured_model = str(st.secrets.get("GEMINI_MODEL", "")).strip()
        retired = {
            "gemini-1.5-flash", "gemini-1.5-pro", "gemini-2.0-flash",
            "gemini-2.0-flash-001", "gemini-2.0-flash-lite",
            "gemini-2.0-flash-lite-001", "gemini-1.5-flash-002",
        }
        # 2026-10 현행 안정 모델. 기존 중단 모델 설정은 자동 대체.
        model = "gemini-3.5-flash-lite" if not configured_model or configured_model in retired else configured_model
        switched = configured_model in retired
        return {"provider": "gemini", "key": gemini_key, "model": model,
                "ready": bool(gemini_key), "switched": switched}

    return {"provider": "openai", "key": openai_key,
            "model": str(st.secrets.get("OPENAI_MODEL", "gpt-4o-mini")).strip() or "gpt-4o-mini",
            "ready": bool(openai_key), "switched": False}


def _safe_ai_error(response, provider):
    """서비스의 원본 오류 본문에는 계정정보가 있을 수 있으므로 그대로 출력하지 않음."""
    status = response.status_code
    if status in (401, 403):
        detail = "API 키가 유효한지, 해당 모델을 사용할 권한이 있는지 확인해 주세요."
    elif status == 404:
        detail = "모델을 찾을 수 없습니다. Streamlit Secrets의 모델명을 확인해 주세요."
    elif status == 429:
        detail = "API 요청 한도 또는 사용량 제한에 걸렸습니다. 사용량/결제 설정을 확인해 주세요."
    elif status == 400:
        detail = "요청 또는 모델 설정이 맞지 않습니다. 모델명을 확인해 주세요."
    elif status >= 500:
        detail = "AI 서비스 쪽 서버 문제가 발생했습니다. 잠시 후 재시도해 주세요."
    else:
        detail = "호출에 실패했습니다. API 설정과 네트워크 상태를 확인해 주세요."
    return f"⚠️ {provider} API 오류 (HTTP {status}): {detail}"


def ask_ai_crm(question, customer_df, orders_db):
    """Gemini/OpenAI 집계 기반 분석. 키 선택, 예외 처리, 결과 상태를 명확하게 함."""
    config = _ai_settings()
    if not config["ready"]:
        return "⚠️ AI 연결 키가 없습니다. Streamlit Secrets에 GEMINI_API_KEY 또는 OPENAI_API_KEY를 설정해 주세요."

    context = make_ai_context(customer_df, orders_db)
    system_prompt = (
        "너는 식혜 온라인 판매자의 CRM 분석 비서다. 한국어로 짧고 실무적으로 답하라. "
        "고객 이름, 전화번호, 주소 등 개인정보는 출력하지 말고, "
        "제공된 집계 수치만 활용하며 근거 없는 정밀 예측을 하지 마라."
    )
    try:
        if config["provider"] == "gemini":
            # 키를 URL 쿼리 파라미터에 넣지 않고 헤더로 전송 (로그 유출 최소화)
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{config['model']}:generateContent"
            headers = {"x-goog-api-key": config["key"], "Content-Type": "application/json"}
            payload = {"contents": [{"parts": [{"text": system_prompt + "\n\n" + context + "\n\n사용자 질문: " + question}]}]}
            res = requests.post(url, headers=headers, json=payload, timeout=55)
            if res.status_code >= 400:
                return _safe_ai_error(res, "Gemini")
            data = res.json()
            chunks = [p.get("text", "") for c in data.get("candidates", [])
                      for p in c.get("content", {}).get("parts", []) if isinstance(p, dict)]
            answer = "\n".join(chunks).strip()
            if answer:
                return answer
            return "⚠️ Gemini가 텍스트를 반환하지 않았습니다. 다시 실행하거나 모델 설정을 확인해 주세요."

        url = "https://api.openai.com/v1/chat/completions"
        headers = {"Authorization": f"Bearer {config['key']}", "Content-Type": "application/json"}
        payload = {"model": config["model"], "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": context + "\n\n질문: " + question}
        ]}
        res = requests.post(url, headers=headers, json=payload, timeout=55)
        if res.status_code >= 400:
            return _safe_ai_error(res, "OpenAI")
        data = res.json()
        return data.get("choices", [{}])[0].get("message", {}).get("content") or "⚠️ AI 응답이 비어 있습니다."
    except requests.exceptions.Timeout:
        return "⚠️ AI 서버 응답 시간이 초과됐습니다. 잠시 후 다시 실행해 주세요."
    except requests.exceptions.RequestException:
        return "⚠️ AI 서버에 연결하지 못했습니다. 네트워크 또는 서비스 상태를 확인해 주세요."
    except (ValueError, KeyError, TypeError, IndexError):
        return "⚠️ AI 응답을 읽는 중 문제가 생겼습니다. 모델 설정을 확인해 주세요."


def ai_status_line():
    config = _ai_settings()
    if not config["ready"]:
        st.warning("AI API 키가 없습니다. Streamlit Secrets에 GEMINI_API_KEY 또는 OPENAI_API_KEY가 필요합니다.")
        return False
    provider_label = "Gemini" if config["provider"] == "gemini" else "OpenAI"
    st.caption(f"연결 설정: {provider_label} · {config['model']} · API 키 등록됨 (실제 호출 성공 여부는 분석 실행 시 확인)")
    if config["switched"]:
        st.info("기존 GEMINI_MODEL에 종료된 모델이 설정돼 있어 지원되는 기본 모델로 자동 전환했습니다.")
    return True



def infer_shipping_item(product_name, channel=""):
    text = str(product_name).strip()
    channel = str(channel)

    if channel == "네이버":
        m_opt = re.search(r'용량\s*:\s*(.*)', text)
        if m_opt:
            text = m_opt.group(1).strip()

    unit_count = 1
    m = re.search(r'(\d+)\s*(개입|개|병|팩|입)', text)
    if m:
        try:
            unit_count = int(m.group(1))
        except Exception:
            unit_count = 1

    if channel == "쿠팡":
        m2 = re.search(r'(?:1\s*L|1리터|1000\s*ml|1000\s*ML).*?(\d+)\s*(개|병|팩)', text, flags=re.I)
        if m2:
            unit_count = int(m2.group(1))
        m3 = re.search(r'(\d+)\s*(개|병|팩).*?(?:1\s*L|1리터|1000\s*ml|1000\s*ML)', text, flags=re.I)
        if m3:
            unit_count = int(m3.group(1))

    if re.search(r'200\s*ML|200\s*ml|200미리|200', text):
        size = "200ml"
    elif re.search(r'500\s*ML|500\s*ml|500미리|500', text):
        size = "500ml"
    elif re.search(r'1\s*L|1리터|1ℓ|1000\s*ML|1000\s*ml', text, flags=re.I):
        size = "1L"
    else:
        size = "기타"

    kind = "단호박식혜" if ("단호박" in text or "호박" in text) else "식혜"
    return f"{kind} {size}", unit_count


def make_shipping_line_items(raw_orders):
    if raw_orders is None or raw_orders.empty:
        return pd.DataFrame(columns=["품목", "채널", "상품명", "주문수량", "개입수", "실제출고수량", "주문일"])

    rows = []
    for _, r in raw_orders.iterrows():
        product = str(r.get("product_name", ""))
        channel = str(r.get("channel", ""))
        qty = int(r.get("quantity", 0)) if pd.notna(r.get("quantity", 0)) else 0
        item, unit_count = infer_shipping_item(product, channel)
        rows.append({
            "품목": item,
            "채널": channel,
            "상품명": product,
            "주문수량": qty,
            "개입수": unit_count,
            "실제출고수량": qty * unit_count,
            "주문일": pd.to_datetime(r.get("order_date", pd.NaT), errors="coerce"),
        })
    return pd.DataFrame(rows)


def make_shipping_summary(raw_orders):
    line = make_shipping_line_items(raw_orders)
    if line.empty:
        return pd.DataFrame(columns=["품목", "실제출고수량", "주문수량", "상품종류수"])

    return (
        line.groupby("품목", dropna=False)
        .agg(
            실제출고수량=("실제출고수량", "sum"),
            주문수량=("주문수량", "sum"),
            상품종류수=("상품명", "nunique"),
        )
        .reset_index()
        .sort_values("실제출고수량", ascending=False)
    )


def make_weekly_shipping_table_from_orders(raw_orders):
    line = make_shipping_line_items(raw_orders)
    if line.empty:
        return pd.DataFrame(), pd.DataFrame()

    line = line.dropna(subset=["주문일"]).copy()
    if line.empty:
        return pd.DataFrame(), pd.DataFrame()

    line["주"] = line["주문일"].dt.to_period("W").astype(str)

    weekly = line.groupby(["주", "품목"], dropna=False)["실제출고수량"].sum().reset_index()
    pivot = weekly.pivot(index="주", columns="품목", values="실제출고수량").fillna(0).astype(int).reset_index()
    avg = weekly.groupby("품목")["실제출고수량"].mean().round(1).reset_index(name="주별평균출고수량")
    avg = avg.sort_values("주별평균출고수량", ascending=False)
    return pivot, avg


def make_weekly_shipping_average(orders_db, weeks=8):
    if orders_db is None or orders_db.empty:
        return pd.DataFrame(columns=["품목", "최근주간평균출고수량", "최근총출고수량", "분석주수"])

    df = orders_db.copy()
    df["order_date"] = pd.to_datetime(df["order_date"], errors="coerce")
    df = df.dropna(subset=["order_date"])
    if df.empty:
        return pd.DataFrame(columns=["품목", "최근주간평균출고수량", "최근총출고수량", "분석주수"])

    max_date = df["order_date"].max()
    start_date = max_date - pd.Timedelta(weeks=weeks)
    df = df[df["order_date"] >= start_date]

    rows = []
    for _, r in df.iterrows():
        product = str(r.get("product_names", ""))
        channel = str(r.get("channel", ""))
        qty = int(r.get("total_quantity", 0)) if pd.notna(r.get("total_quantity", 0)) else 0
        item, unit_count = infer_shipping_item(product, channel)
        rows.append({"품목": item, "실제출고수량": qty * unit_count})

    out = pd.DataFrame(rows)
    if out.empty:
        return pd.DataFrame(columns=["품목", "최근주간평균출고수량", "최근총출고수량", "분석주수"])

    summary = out.groupby("품목", dropna=False).agg(최근총출고수량=("실제출고수량", "sum")).reset_index()
    summary["분석주수"] = weeks
    summary["최근주간평균출고수량"] = (summary["최근총출고수량"] / weeks).round(1)
    return summary[["품목", "최근주간평균출고수량", "최근총출고수량", "분석주수"]].sort_values("최근주간평균출고수량", ascending=False)


# 같은 파일을 다시 렌더링할 때 복호화/엑셀 분석을 반복하지 않습니다.
@st.cache_data(ttl=900, max_entries=12, show_spinner=False)
def parse_uploaded_file_cached(filename, content):
    file_like = io.BytesIO(content)
    if filename.lower().endswith(".csv"):
        # 쿠팡/네이버 CSV는 UTF-8과 CP949가 섞일 수 있음.
        raw = None
        for encoding in ("utf-8-sig", "cp949", "euc-kr"):
            try:
                raw = pd.read_csv(io.BytesIO(content), encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
        if raw is None:
            raise ValueError("CSV 인코딩을 판별하지 못했습니다. UTF-8 또는 CP949로 저장해 주세요.")
    else:
        raw = read_excel_file(file_like)
    market = detect_market(filename, raw)
    if market == "네이버":
        necessary = ["상품주문번호", "수취인명", "수량", "결제일"]
        missing = [name for name in necessary if name not in raw.columns]
        if missing:
            raise ValueError(
                "이 네이버 완성 파일은 주문 분석용 원본 정보가 없는 구형 발송처리 파일입니다. "
                "고객/상품 데이터를 복원할 수 없어 분석할 수 없습니다. "
                "송장 업무 탭에서 원본 4개 파일로 다시 생성해 주세요. "
                "부족한 열: " + ", ".join(missing)
            )
    std = standardize(raw, market, filename)
    return std, market, len(raw)


@st.cache_data(ttl=900, max_entries=5, show_spinner=False)
def _prepare_uploaded_files(file_blobs):
    """파일들의 바이트가 같으면 엑셀 파싱과 주문 단위 병합을 반복하지 않는다."""
    import hashlib
    frames, errors, logs = [], [], []
    seen_digests = set()
    for filename, content in file_blobs:
        # V30 이후 ★네이버 완성 파일은 주문 원본 72열도 포함해 CRM 분석 가능.
        # 구버전 V28/V29에서 이미 만든 4열 전용 파일은 고객정보가 없어 복구 불가능.
        short_name = str(filename).strip()
        digest = hashlib.sha256(content).digest()
        if digest in seen_digests:
            logs.append({"파일명": filename, "채널": "", "원본행": 0,
                         "정리행": 0, "상태": "중복 파일 제외"})
            continue
        seen_digests.add(digest)
        try:
            std, market, raw_len = parse_uploaded_file_cached(filename, content)
            frames.append(std)
            logs.append({"파일명": filename, "채널": market, "원본행": raw_len,
                         "정리행": len(std), "상태": "성공"})
        except Exception as exc:
            errors.append((filename, str(exc)))
            logs.append({"파일명": filename, "채널": "", "원본행": 0,
                         "정리행": 0, "상태": "실패"})
    if not frames:
        return pd.DataFrame(), pd.DataFrame(), logs, errors
    raw = pd.concat(frames, ignore_index=True)
    # 네이버 상품주문번호는 상품행 고유값이라 겹친 엑셀에서 다시 나오면 중복 출고다.
    # 쿠팡 주문번호는 여러 상품행에 재사용되므로 여기에 같은 중복 제거를 적용하면 안 된다.
    is_naver = raw["channel"].eq("네이버")
    product_id = raw["product_order_no"].fillna("").astype(str).str.strip()
    has_product_id = ~product_id.isin(("", "nan", "None", "none"))
    repeated = is_naver & has_product_id & raw.duplicated(subset=["channel", "product_order_no"], keep="first")
    if repeated.any():
        skipped = int(repeated.sum())
        raw = raw.loc[~repeated].reset_index(drop=True)
        logs.append({"파일명": "여러 파일의 중복 상품주문번호", "채널": "네이버", "원본행": skipped,
                     "정리행": 0, "상태": "중복 상품행 제외"})
    order_level = make_order_level(raw, use_strict=False)
    return raw, order_level, logs, errors


# ===== 재고관리 통합 기능 (별도 모듈 필요 없음) =====
# 품목 코드는 저장 후 이름이 바뀌더라도 유지되는 키로 사용한다.
ITEMS = [
    ("sikhye_1l", "식혜 1L", 150, False),
    ("sikhye_500ml", "식혜 500ml", 100, True),
    ("sikhye_200ml", "식혜 200ml", 100, False),
    ("pumpkin_1l", "단호박식혜 1L", 60, False),
    ("pumpkin_500ml", "단호박식혜 500ml", 30, False),
]
KST = ZoneInfo("Asia/Seoul")
KEYS = [p[0] for p in ITEMS]
LABELS = {p[0]: p[1] for p in ITEMS}


def _safe_int(x):
    try:
        return max(0, int(x))
    except (ValueError, TypeError, OverflowError):
        return 0


def _is_low(key, remaining):
    limit, inclusive = next((p[2], p[3]) for p in ITEMS if p[0] == key)
    return remaining <= limit if inclusive else remaining < limit


def classify_label(label):
    s = str(label).lower().replace(" ", "")
    pumpkin = "단호박" in s or "호박" in s
    if "200ml" in s:
        return None if pumpkin else "sikhye_200ml"
    if "500ml" in s:
        return "pumpkin_500ml" if pumpkin else "sikhye_500ml"
    if "1l" in s or "1000ml" in s:
        return "pumpkin_1l" if pumpkin else "sikhye_1l"
    return None


def today_shipment(raw_orders, shipping_summary_func):
    """기존 쿠팡/네이버 수량 파서를 재사용해 출고량을 정확히 가져온다."""
    out = {key: 0 for key in KEYS}
    unknown = []
    if raw_orders is None or raw_orders.empty:
        return out, unknown
    shipping = shipping_summary_func(raw_orders)
    for _, r in shipping.iterrows():
        key = classify_label(r.get("품목", ""))
        if key:
            out[key] += _safe_int(r.get("실제출고수량"))
        else:
            unknown.append((str(r.get("품목", "기타")), _safe_int(r.get("실제출고수량"))))
    return out, unknown


def _get_db_client():
    """공개용 publishable key로 재고를 수정할 수 없도록 별도의 서버 비밀키 사용."""
    secret = st.secrets.get("INVENTORY_DB_KEY", "")
    url = st.secrets.get("SUPABASE_URL", "")
    if not secret or not url:
        return None
    if not (secret.startswith("sb_secret_") or secret.startswith("eyJ")):
        return None
    try:
        from supabase import create_client
        return create_client(url, secret)
    except Exception:
        return None


@st.cache_data(ttl=90, show_spinner=False)
def _read_inventory_data(url, secret, nonce=0):
    """캐시 키에는 비밀키가 포함되지만 UI에 출력하지 않는다."""
    from supabase import create_client
    client = create_client(url, secret)
    rows = []
    start = 0
    while True:
        result = client.table("inventory_daily").select("inventory_date,item_key,opening_qty,received_qty,auto_ship_qty,actual_ship_qty,closing_qty,note,updated_at").order("inventory_date", desc=True).order("item_key").range(start, start + 999).execute()
        batch = result.data or []
        rows.extend(batch)
        if len(batch) < 1000:
            break
        start += 1000
    return rows


def fetch_inventory_rows():
    secret = st.secrets.get("INVENTORY_DB_KEY", "")
    url = st.secrets.get("SUPABASE_URL", "")
    if not secret or not url:
        return [], "DB 저장 준비 중: Streamlit Secrets에 INVENTORY_DB_KEY를 추가해주세요."
    try:
        return _read_inventory_data(url, secret), ""
    except Exception as exc:
        # DB 테이블 미생성, 네트워크 오류 등을 앱 전체 장애로 확대하지 않는다.
        return [], f"재고 DB를 불러오지 못했습니다. inventory_schema.sql 적용 여부와 DB 연결을 확인하세요. ({type(exc).__name__})"


def previous_inventory(rows, as_of):
    prev = {}
    for r in rows:
        try:
            d = date.fromisoformat(str(r["inventory_date"])[:10])
            key = r["item_key"]
            if d < as_of and key in KEYS and (key not in prev or d > prev[key][0]):
                prev[key] = (d, _safe_int(r["closing_qty"]))
        except (KeyError, ValueError, TypeError):
            pass
    return prev


def existing_today(rows, as_of):
    return {
        str(r.get("item_key")): r
        for r in rows if str(r.get("inventory_date"))[:10] == as_of.isoformat()
    }


# 카카오톡에서 글꼴 폭이 달라도 보기 좋도록 공백 정렬 대신 한 줄씩 구성.
# 칠판에 옮겨 적기 쉽도록 단위 '개', 긴 품목명, 요약 경고를 제외한다.
COPY_NAMES = {
    "sikhye_1l": "1L",
    "sikhye_500ml": "500ML",
    "sikhye_200ml": "200ML",
    "pumpkin_1l": "단호박 1L",
    "pumpkin_500ml": "단호박 500ML",
}


def make_message(as_of, rows, compact=False):
    header = f"📦 재고 ({as_of:%m/%d})" if compact else f"📦 출고 / 재고 ({as_of:%m/%d})"
    lines = [header, ""]
    # 저장된 records에 같은 품목이 없으면 생략하고, 품목 순서는 고정한다.
    by_key = {r["item_key"]: r for r in rows}
    for key in KEYS:
        r = by_key.get(key)
        if r is None:
            continue
        name = COPY_NAMES[key]
        warning = "   ⚠️" if _is_low(key, r["closing_qty"]) else ""
        if compact:
            lines.append(f"{name}: {r['closing_qty']:,}{warning}")
        else:
            lines.append(f"{name}: 출고 {r['actual_ship_qty']:,} / 재고 {r['closing_qty']:,}{warning}")
    return "\n".join(lines)


def render_oneclick_copy_button(message, label, button_id, accent=False):
    """한 번 클릭하면 실제 클립보드에 복사 (Streamlit 1.56 st.html JS 지원).

    복사 텍스트는 숫자/품목 이름으로 만든 값을 JSON으로 직렬화하며
    HTML 태그나 스크립트로 해석되지 않도록 특수문자를 이스케이프한다.
    """
    # 비신뢰 데이터가 JS 문자열 경계를 벗어나지 않도록 < 등 이스케이프.
    payload = json.dumps(message, ensure_ascii=False).replace('<', r'\u003c').replace('>', r'\u003e').replace('&', r'\u0026')
    button_text = json.dumps(label, ensure_ascii=False)
    tone = '#eaf7ee' if accent else '#f7f8fa'
    border = '#89c7a0' if accent else '#c9cdd3'
    markup = f"""
    <div class="inventory-copy-root" id="{button_id}">
      <button class="inventory-copy-button" type="button" aria-label={label}>
        <span class="copy-label"></span>
      </button>
      <span class="copy-status" aria-live="polite"></span>
    </div>
    <style>
      #{button_id} {{display:flex; align-items:center; gap:6px; width:100%; font-family:inherit;}}
      #{button_id} .inventory-copy-button {{cursor:pointer; font-family:inherit; font-size:15px; font-weight:700; line-height:1.4;
        box-sizing:border-box; width:100%; min-height:44px; padding:9px 9px; border-radius:10px;
        background:{tone}; color:#173023; border:1px solid {border}; white-space:nowrap;}}
      #{button_id} .inventory-copy-button:hover {{filter:brightness(.97)}}
      #{button_id} .inventory-copy-button:focus-visible {{outline:2px solid #2a8a5b; outline-offset:2px}}
      #{button_id} .copy-status {{font-size:11px; white-space:nowrap; color:#ad2222;}}
    </style>
    <script>
      (() => {{
        const root = document.getElementById({json.dumps(button_id)});
        if (!root) return;
        const button = root.querySelector('button');
        const label = root.querySelector('.copy-label');
        const status = root.querySelector('.copy-status');
        const originalLabel = {button_text};
        const copyValue = {payload};
        label.textContent = originalLabel;
        // Streamlit 재렌더링 시 최신 문구로 이벤트를 덮어쓴다.
        let resetTimer;
        function fallbackCopy(text) {{
          const node = document.createElement('textarea');
          node.value = text;
          node.setAttribute('readonly', '');
          node.style.cssText = 'position:fixed;left:-9999px;top:0;opacity:0;';
          document.body.appendChild(node);
          node.focus();
          node.select();
          try {{ return document.execCommand('copy'); }}
          finally {{ node.remove(); }}
        }}
        button.onclick = async () => {{
          let ok = false;
          try {{
            if (navigator.clipboard && window.isSecureContext) {{
              await navigator.clipboard.writeText(copyValue);
              ok = true;
            }} else {{ ok = fallbackCopy(copyValue); }}
          }} catch (error) {{
            try {{ ok = fallbackCopy(copyValue); }} catch (_) {{ ok = false; }}
          }}
          clearTimeout(resetTimer);
          label.textContent = ok ? '✅ 복사 완료!' : originalLabel;
          status.textContent = ok ? '' : '복사 불가';
          resetTimer = setTimeout(() => {{ label.textContent = originalLabel; status.textContent = ''; }}, 1500);
        }};
      }})();
    </script>
    """
    st.html(markup, width="stretch", unsafe_allow_javascript=True)


def snapshots_weekly(rows, weeks=12, as_of=None):
    """날짜별 스냅샷을 동일 주·품목으로 합산. 기록한 날짜 수 함께 제공."""
    if as_of is None:
        as_of = datetime.now(KST).date()
    current_week = as_of - timedelta(days=as_of.weekday())
    start_week = current_week - timedelta(weeks=weeks - 1)
    data = []
    for r in rows:
        try:
            d = date.fromisoformat(str(r.get("inventory_date", ""))[:10])
            if start_week <= d <= as_of and d.weekday() < 5 and r.get("item_key") in KEYS:
                start = d - timedelta(days=d.weekday())
                data.append({"주": start.isoformat(), "품목": LABELS[r["item_key"]], "출고수량": _safe_int(r.get("actual_ship_qty")), "날짜": d.isoformat()})
        except ValueError:
            continue
    if not data:
        return pd.DataFrame(), pd.DataFrame()
    df = pd.DataFrame(data)
    weekly = df.groupby(["주", "품목"], as_index=False).agg(출고수량=("출고수량", "sum"), 기록일수=("날짜", "nunique"))
    # 평일 5일 기록을 모두 남긴 주만 평균 산출한다. 미기록일은 0출고로 취급하지 않는다.
    completed = weekly[weekly["기록일수"] >= 5]
    avg = (completed.groupby("품목")["출고수량"].agg(["mean", "count"])
           .reset_index().rename(columns={"mean": "완전한_주간_평균", "count": "분석주수"})) if not completed.empty else pd.DataFrame()
    return weekly, avg


def _order_dates_kst(orders_db):
    """주문 시각을 한국 날짜로 통일한다. 기존 주문 테이블은 UTC ISO 문자열을 저장한다."""
    return pd.to_datetime(orders_db["order_date"], utc=True, errors="coerce").dt.tz_convert(KST).dt.date


def historical_weekday_report(orders_db, as_of, infer_fn):
    """작년 같은 ISO주 월~금: 주문 건수(정확)와 단일 상품 주문의 식별 가능한 병수(부분 집계).

    비교 평균: 동일 주를 포함한 작년 직전 4주간 월~금의 평균. 주문 기록이
    존재하는 주만 평균 대상으로 하며, 없는 주를 0으로 취급하지 않는다.
    """
    iso = as_of.isocalendar()
    try:
        monday = date.fromisocalendar(iso.year - 1, iso.week, 1)
    except ValueError:
        # ISO 53주가 직전 연도에 없다면 52주로 조정하고 화면에 기간을 명시한다.
        monday = date.fromisocalendar(iso.year - 1, 52, 1)
    days = [monday + timedelta(days=i) for i in range(5)]
    week_starts = [monday - timedelta(weeks=i) for i in range(3, -1, -1)]
    empty_day = {d: {"orders": 0, "eligible": 0, "qty": {k: 0 for k in KEYS}} for d in days}
    result = {
        "days": days, "monday": monday, "daily": empty_day,
        "weekly_4": {}, "4week_coverage": 0,
        "total_orders": 0, "eligible": 0,
        "qty": {k: 0 for k in KEYS}, "average_orders_4w": None,
        "average_qty_4w": {k: None for k in KEYS},
    }
    if orders_db is None or orders_db.empty or "order_date" not in orders_db.columns:
        return result
    df = orders_db.copy()
    df["_local_date"] = _order_dates_kst(df)
    # 전체 수년치 주문에서 날짜별 상품을 모두 분석하지 말고 필요한 4주치만 처리한다.
    first_day = week_starts[0]
    last_day = days[-1]
    df = df[df["_local_date"].between(first_day, last_day)].copy()
    if df.empty:
        return result
    # 하나의 주문이 하나의 원본 상품행인 경우에 한해 품목과 구성 병수를 추정한다.
    df["_single"] = pd.to_numeric(df.get("raw_rows", pd.Series(index=df.index, dtype="float64")), errors="coerce").eq(1)
    df["_qty"] = pd.to_numeric(df.get("total_quantity", pd.Series(index=df.index, dtype="float64")), errors="coerce").fillna(0)
    df["_item"] = None
    df["_units"] = 0
    valid = df["_single"] & df["_qty"].gt(0)
    for idx, row in df.loc[valid].iterrows():
        name = str(row.get("product_names", ""))
        if not name or name.lower() == "nan":
            continue
        item_name, pack_units = infer_fn(name, str(row.get("channel", "")))
        key = classify_label(item_name)
        if key and _safe_int(pack_units) > 0:
            df.at[idx, "_item"] = key
            df.at[idx, "_units"] = _safe_int(pack_units)
    # 데이터를 갖고 있는 4주만 계산. 미기록은 영업일 0출고가 아니다.
    for monday_i in week_starts:
        weekly_days = [monday_i + timedelta(days=i) for i in range(5)]
        frame = df[df["_local_date"].isin(weekly_days)]
        if frame.empty:
            continue
        counts = {k: 0 for k in KEYS}
        recognized = frame[frame["_item"].isin(KEYS)]
        for _, row in recognized.iterrows():
            counts[row["_item"]] += _safe_int(row["_qty"] * row["_units"])
        result["weekly_4"][monday_i] = {"orders": len(frame), "eligible": len(recognized), "qty": counts}
    result["4week_coverage"] = len(result["weekly_4"])
    if result["4week_coverage"]:
        n = result["4week_coverage"]
        result["average_orders_4w"] = sum(x["orders"] for x in result["weekly_4"].values()) / n
        for key in KEYS:
            result["average_qty_4w"][key] = sum(x["qty"][key] for x in result["weekly_4"].values()) / n
    for d in days:
        frame = df[df["_local_date"] == d]
        result["daily"][d]["orders"] = len(frame)
        recognized = frame[frame["_item"].isin(KEYS)]
        result["daily"][d]["eligible"] = len(recognized)
        for _, row in recognized.iterrows():
            key = row["_item"]
            q = _safe_int(row["_qty"] * row["_units"])
            result["daily"][d]["qty"][key] += q
            result["qty"][key] += q
        result["total_orders"] += len(frame)
        result["eligible"] += len(recognized)
    return result


def render_historical_production_report(orders_db, as_of, infer_fn):
    """작년 월~금 주문/품목별 식별 가능한 수량과 주간 평균을 바로 보여준다."""
    report = historical_weekday_report(orders_db, as_of, infer_fn)
    days = report["days"]
    monday = report["monday"]
    friday = days[-1]
    st.markdown("### 📅 작년 같은 주 월~금 출고 비교")
    st.caption(f"대상 기간: **{monday:%Y.%m.%d}(월) ~ {friday:%m.%d}(금)**  |  주문일 기준  |  날짜별·품목별 비교")
    d_cols = st.columns(3)
    d_cols[0].metric("작년 월~금 주문", f"{report['total_orders']:,}건")
    d_cols[1].metric("작년 하루 평균 주문", f"{report['total_orders']/5:,.1f}건")
    avg_orders = report["average_orders_4w"]
    d_cols[2].metric("작년 인접 4주 월~금 평균", f"{avg_orders:,.1f}건" if avg_orders is not None else "자료 없음")
    if report["total_orders"] == 0:
        st.info("작년 해당 월~금에 저장된 주문이 없습니다. DB에 그 기간 주문이 들어있는지 확인해주세요. 기록이 없다고 실제 출고가 0개였던 것은 아닙니다.")
    n_weeks = report["4week_coverage"]
    if n_weeks < 4:
        st.caption(f"작년 같은 주와 직전 3주 중 주문 기록이 확인되는 기간: **{n_weeks}/4주**. 미기록 주를 0으로 처리하지 않습니다.")
    st.markdown("**품목별 월~금 수량 비교**")
    labels = ["월", "화", "수", "목", "금"]
    rows = []
    for key, label, _, _ in ITEMS:
        avg = report["average_qty_4w"][key]
        row = {"품목": label}
        for d, day_label in zip(days, labels):
            row[f"{day_label} {d:%m/%d}"] = report["daily"][d]["qty"][key] if report["total_orders"] else "—"
        row["5일 합계"] = report["qty"][key] if report["total_orders"] else "—"
        row["하루 평균"] = round(report["qty"][key] / 5, 1) if report["total_orders"] else "—"
        row["인접 주간 평균"] = round(avg, 1) if avg is not None else "—"
        rows.append(row)
    st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)
    st.caption(
        f"⚠️ 위의 품목 수량은 **단일상품 주문 {report['eligible']:,}건 / 해당 주 전체 {report['total_orders']:,}건**을 이용해 계산한 "
        "**확인 가능한 부분 수량**입니다. 혼합 주문과 택배 외 출고는 누락될 수 있어 실제 전체 출고량과 다릅니다. "
        f"'인접 주간 평균'은 작년 해당 주와 그 이전 3주 중 기록이 있는 {n_weeks}주의 월~금 평균입니다. "
        "정확한 생산량 판단에는 과거 원본 상품별 주문 엑셀이 필요합니다."
    )
    if report["total_orders"] > 0:
        day_orders = "  |  ".join(
            f"{weekday} {report['daily'][day]['orders']:,}건"
            for day, weekday in zip(days, labels)
        )
        st.caption("📋 작년 요일별 주문 건수: " + day_orders)


def _recorded_inventory_calendar(rows):
    """확정 저장된 실출고 스냅샷만 날짜별로 묶는다. 미기록일은 0과 구분한다."""
    calendar = {}
    for row in rows:
        try:
            day = date.fromisoformat(str(row.get("inventory_date", ""))[:10])
            key = row.get("item_key")
            if key not in KEYS:
                continue
            calendar.setdefault(day, {})[key] = _safe_int(row.get("actual_ship_qty"))
        except (ValueError, TypeError):
            pass
    return calendar


def _week_snapshot(calendar, monday):
    """품목별 월~금 실출고 및 기록 현황. 미기록일은 None; 0은 실제 0개."""
    days = [monday + timedelta(days=i) for i in range(5)]
    complete_days = [d for d in days if all(k in calendar.get(d, {}) for k in KEYS)]
    result = {}
    for key in KEYS:
        quantities = [calendar.get(d, {}).get(key) for d in days]
        result[key] = {
            "daily": quantities,
            "recorded": sum(q is not None for q in quantities),
            "partial": sum(q for q in quantities if q is not None),
            "total": sum(quantities) if all(q is not None for q in quantities) else None,
        }
    return {"days": days, "complete_days": len(complete_days), "items": result}


def _prior_year_monday(as_of):
    iso = as_of.isocalendar()
    try:
        return date.fromisocalendar(iso.year - 1, iso.week, 1)
    except ValueError:
        # ISO 53주가 전년에 없다면 52주를 참조한다.
        return date.fromisocalendar(iso.year - 1, 52, 1)


def _completed_week_average(calendar, anchor_monday, count=4):
    """기준 주 이전 count주만 비교. 5일 모두 기록된 주만 평균에 포함."""
    completed = []
    for i in range(count, 0, -1):
        monday = anchor_monday - timedelta(weeks=i)
        week = _week_snapshot(calendar, monday)
        if week["complete_days"] == 5:
            completed.append(week)
    if not completed:
        return {k: None for k in KEYS}, 0
    return {
        k: sum(w["items"][k]["total"] for w in completed) / len(completed)
        for k in KEYS
    }, len(completed)


def render_verified_inventory_history(history_rows, as_of):
    """내년에는 작년 같은 ISO주 실제 월~금 기록으로 바뀌는 비교 영역."""
    calendar = _recorded_inventory_calendar(history_rows)
    recorded_days = sorted(d for d, values in calendar.items() if all(k in values for k in KEYS))
    this_monday = as_of - timedelta(days=as_of.weekday())
    last_monday = _prior_year_monday(as_of)
    current = _week_snapshot(calendar, this_monday)
    prior = _week_snapshot(calendar, last_monday)
    prev_avg, prev_count = _completed_week_average(calendar, this_monday, count=4)
    last_avg, last_count = _completed_week_average(calendar, last_monday, count=4)

    st.markdown("### 📆 실제 출고 기록 · 주간 생산 참고")
    st.caption("주문 DB 추정치가 아니라, **재고관리에서 최종 저장한 '오늘 출고'** 수량만 사용합니다. 월~금 기준입니다.")
    cols = st.columns(3)
    cols[0].metric("출고 기록한 날짜", f"{len(recorded_days):,}일")
    cols[1].metric("이번 주 기록", f"{current['complete_days']}/5일")
    cols[2].metric("최근 4주 월~금 평균", f"{prev_count}/4주 확보")
    if recorded_days:
        st.caption(f"기록 시작 **{recorded_days[0]:%Y.%m.%d}** · 마지막 저장 **{recorded_days[-1]:%Y.%m.%d}** · 미기록일은 출고 0개로 처리하지 않습니다.")
    else:
        st.info("아직 확정된 출고 기록이 없습니다. 위에서 매일 수량을 확인한 뒤 **'오늘 재고 최종 저장'**을 누르면 오늘부터 기록이 쌓입니다. 출고가 0개인 날도 저장해야 정확한 주간 평균을 계산할 수 있습니다.")

    st.markdown(f"**이번 주 | {this_monday:%Y.%m.%d}(월) ~ {(this_monday + timedelta(days=4)):%m.%d}(금)**")
    cols = [f"{wk} {day:%m/%d}" for wk, day in zip(["월","화","수","목","금"], current["days"])]
    table_rows = []
    for key, name, _, _ in ITEMS:
        info = current["items"][key]
        row = {"품목": name}
        row.update({c: (f"{v:,}" if v is not None else "—") for c, v in zip(cols, info["daily"])})
        row["기록된 합계"] = f"{info['partial']:,}" if info["recorded"] else "—"
        row["이전 4주 평균"] = f"{prev_avg[key]:,.1f}" if prev_avg[key] is not None else "—"
        table_rows.append(row)
    st.dataframe(pd.DataFrame(table_rows), use_container_width=True, hide_index=True)
    st.caption("'기록된 합계'는 저장된 날짜만 합산합니다. **— = 미기록**, **0 = 0개 출고를 확인하고 저장**한 날입니다. 4주 평균은 월~금 모두 기록된 주만 계산합니다.")

    st.markdown(f"**작년 같은 주 | {last_monday:%Y.%m.%d}(월) ~ {(last_monday + timedelta(days=4)):%m.%d}(금)**")
    if prior["complete_days"] == 0:
        st.info("작년 같은 주의 **확정된 실제 출고 기록이 아직 없습니다.** 과거 주문 DB의 부분 집계는 생산량으로 오인하기 쉬워 여기에는 표시하지 않습니다. 오늘부터 저장하면 내년 같은 주 비교에 자동으로 표시돼요.")
    else:
        st.caption(f"작년 기록 {prior['complete_days']}/5일 · 작년 직전 4주 중 완전 기록 {last_count}/4주")
        prior_cols = [f"{wk} {day:%m/%d}" for wk, day in zip(["월","화","수","목","금"], prior["days"])]
        history_table = []
        for key, name, _, _ in ITEMS:
            info = prior["items"][key]
            row = {"품목": name}
            row.update({c: (f"{v:,}" if v is not None else "—") for c, v in zip(prior_cols, info["daily"])})
            row["월~금 합계"] = f"{info['total']:,}" if info["total"] is not None else "—"
            row["하루 평균"] = f"{info['total']/5:,.1f}" if info["total"] is not None else "—"
            row["작년 직전 4주 평균"] = f"{last_avg[key]:,.1f}" if last_avg[key] is not None else "—"
            history_table.append(row)
        st.dataframe(pd.DataFrame(history_table), use_container_width=True, hide_index=True)
        if prior["complete_days"] < 5:
            st.warning("작년 해당 주의 일부 날짜가 빠져 있어 **주간 합계와 일평균은 표시하지 않습니다.** 빠진 날짜를 0개로 간주하면 실제보다 출고량이 작아집니다.")
    st.caption("같은 ISO 주차의 월~금을 비교하므로 해마다 날짜는 조금씩 다를 수 있습니다. 명절이나 배송 휴무일은 별도로 고려해 생산량을 결정하세요.")


def save_inventory(records, summary_note="", expected_versions=None):
    client = _get_db_client()
    if client is None:
        raise RuntimeError("재고 DB 비밀키를 설정해야 저장할 수 있습니다.")
    # 두 화면에서 같은 날짜를 열어놓고 저장할 때 조용히 덮어쓰지 않도록 점검.
    # 완전한 원자적 잠금은 아니지만, 실수로 과거 스냅샷을 저장하는 일을 줄인다.
    if expected_versions is not None and records:
        save_day = str(records[0]["inventory_date"])
        current_data = client.table("inventory_daily").select("item_key,updated_at").eq("inventory_date", save_day).execute()
        live_versions = {str(row["item_key"]): str(row.get("updated_at", ""))
                         for row in (current_data.data or [])}
        expected = {str(key): str(version) for key, version in expected_versions.items()}
        if any(live_versions.get(key, "") != expected.get(key, "") for key in KEYS):
            raise RuntimeError("다른 화면에서 오늘 재고가 변경됐습니다. 새로고침해 최신 수량을 확인한 다음 저장해 주세요.")
    # 재고 재저장 시 같은 날짜·품목 행은 업데이트된다. 이전 재고는 추가로 차감되지 않는다.
    payload = []
    for r in records:
        payload.append({
            "inventory_date": str(r["inventory_date"]),
            "item_key": r["item_key"],
            "opening_qty": _safe_int(r["opening_qty"]),
            "received_qty": _safe_int(r["received_qty"]),
            "auto_ship_qty": _safe_int(r["auto_ship_qty"]),
            "actual_ship_qty": _safe_int(r["actual_ship_qty"]),
            "closing_qty": _safe_int(r["closing_qty"]),
            "note": str(summary_note)[:300],
            "updated_at": datetime.now(KST).isoformat(),
        })
    client.table("inventory_daily").upsert(payload, on_conflict="inventory_date,item_key").execute()
    _read_inventory_data.clear()
    # audit는 스냅샷이 성공적으로 저장된 경우에만 추가한다.
    try:
        audit = [{"inventory_date": p["inventory_date"], "item_key": p["item_key"],
                  "opening_qty": p["opening_qty"], "received_qty": p["received_qty"],
                  "auto_ship_qty": p["auto_ship_qty"], "actual_ship_qty": p["actual_ship_qty"],
                  "closing_qty": p["closing_qty"], "note": p["note"]} for p in payload]
        client.table("inventory_audit").insert(audit).execute()
        return True
    except Exception:
        return False


def _draft_key(as_of, key, field):
    return f"inventory_{as_of.isoformat()}_{key}_{field}"


def _weekly_order_counts(orders_db, as_of):
    """주문 DB의 날짜로 ISO 연도/주를 비교한다. 병 수량이 아닌 주문 건수."""
    this_year, this_week, _ = as_of.isocalendar()
    if orders_db is None or orders_db.empty or "order_date" not in orders_db.columns:
        return 0, 0
    dates = pd.to_datetime(orders_db["order_date"], utc=True, errors="coerce")
    weeks = dates.dt.isocalendar()
    cur = int(((weeks["year"] == this_year) & (weeks["week"] == this_week)).sum())
    prior = int(((weeks["year"] == this_year - 1) & (weeks["week"] == this_week)).sum())
    return cur, prior


def render_inventory(raw_orders, orders_db, shipping_summary_func, infer_shipping_item):
    as_of = datetime.now(KST).date()
    # 맨 위에 복사 버튼 2개 + 미리보기 1개 배치.
    # 입력값은 아래에서 계산되므로 자리만 먼저 확보한 후 채운다.
    title_col, copy_all_col, copy_stock_col, preview_col = st.columns(
        [4.1, 2.25, 2.15, 1.5], gap="small", vertical_alignment="center"
    )
    with title_col:
        st.subheader("📦 재고관리")
    with copy_all_col:
        copy_all_slot = st.empty()
    with copy_stock_col:
        copy_stock_slot = st.empty()
    with preview_col:
        preview_slot = st.empty()
    st.caption(f"{as_of:%Y년 %m월 %d일} 기준  |  출고와 기존 재고 확인 → 저장  |  저장한 실출고량이 내년 같은 주 생산 참고자료가 됩니다.")

    history_rows, error = fetch_inventory_rows()
    today_saved = existing_today(history_rows, as_of)
    prev_saved = previous_inventory(history_rows, as_of)
    auto, unknown = today_shipment(raw_orders, shipping_summary_func)
    if error:
        st.warning(error)
    elif len(today_saved) == len(KEYS):
        st.success("✅ 오늘 실출고 5개 품목이 DB에 저장되어 있습니다. 수량을 바꾸면 다시 **최종 저장**해 덮어쓸 수 있어요.")
    else:
        st.info("📝 오늘 실출고는 아직 미저장 상태입니다. 엑셀과 수동 출고를 확인한 뒤 **최종 저장**을 눌러주세요.")
    if raw_orders is None or raw_orders.empty:
        st.caption("ℹ️ 오늘 업로드한 엑셀이 없어 자동 출고가 0개입니다. 이전 저장값이 있으면 그대로 표시됩니다.")
    if unknown:
        st.warning("자동 분류되지 않은 품목: " + ", ".join(f"{a} {b}개" for a, b in unknown))

    st.markdown("### 오늘 품목별 재고")
    st.caption("**기존 재고** = 오늘 출고 전에 남아 있는 수량(생산 후에는 이 수량을 직접 수정)  /  **오늘 출고** = 실제 나간 수량  /  **남은 재고** = 자동 계산")

    head = st.columns([2.2, 1.55, 1.55, 1.4], gap="small")
    for col, label in zip(head, ["품목", "오늘 출고 ✏️", "기존 재고 ✏️", "남은 재고"]):
        col.markdown(f"**{label}**")

    records = []
    for key, label, threshold, inclusive in ITEMS:
        saved = today_saved.get(key, {})
        # 예전 버전의 생산/입고값도 기존 재고에 합쳐서 불러와야 재고가 줄어들지 않는다.
        opening_default = (
            _safe_int(saved.get("opening_qty", 0)) + _safe_int(saved.get("received_qty", 0))
            if saved else prev_saved.get(key, (None, 0))[1]
        )
        actual_default = _safe_int(saved.get("actual_ship_qty", auto[key]))
        actual_key = _draft_key(as_of, key, "actual")
        auto_key = _draft_key(as_of, key, "last_auto")
        opening_key = _draft_key(as_of, key, "stock_v13")
        # Streamlit의 지연 렌더링에서 탭을 옮기면 위젯 키가 정리될 수 있다.
        # 확정 저장 전의 수량은 별도 세션 초안 키에 보관해 다시 돌아왔을 때 복구한다.
        draft_actual_key = f"_draft_keep_{actual_key}"
        draft_opening_key = f"_draft_keep_{opening_key}"
        opening_default = int(st.session_state.get(draft_opening_key, opening_default))
        if actual_key not in st.session_state:
            last_auto = int(st.session_state.get(auto_key, auto[key]))
            retained = int(st.session_state.get(draft_actual_key, actual_default))
            st.session_state[actual_key] = max(0, retained + int(auto[key]) - last_auto)
            st.session_state[auto_key] = auto[key]
        elif st.session_state.get(auto_key, auto[key]) != auto[key]:
            # 엑셀이 추가/변경돼도 기존 직접 수정분은 보존한다.
            delta = auto[key] - st.session_state[auto_key]
            st.session_state[actual_key] = max(0, int(st.session_state[actual_key]) + delta)
            st.session_state[auto_key] = auto[key]

        with st.container(border=True):
            c_name, c_out, c_stock, c_left = st.columns([2.2, 1.55, 1.55, 1.4], gap="small", vertical_alignment="center")
            with c_name:
                st.markdown(f"**{label}**")
                st.caption(f"엑셀 자동 {auto[key]:,}개  ·  부족 기준 {threshold}개{' 이하' if inclusive else ' 미만'}")
            with c_out:
                actual = st.number_input(
                    f"{label} 오늘 출고", min_value=0, step=1,
                    key=actual_key, label_visibility="collapsed",
                )
            with c_stock:
                opening = st.number_input(
                    f"{label} 기존 재고", min_value=0, step=1,
                    value=opening_default, key=opening_key, label_visibility="collapsed",
                )
            # 위젯이 표시된 현재 입력값을 별도 초안으로 유지 (DB 저장 아님).
            st.session_state[draft_actual_key] = int(actual)
            st.session_state[draft_opening_key] = int(opening)
            closing = int(opening) - int(actual)
            low = _is_low(key, closing)
            with c_left:
                color = "#c62828" if low or closing < 0 else "#18794e"
                status = "🔴 부족" if low else "🟢 정상"
                st.markdown(
                    f'<div style="font-size:23px;font-weight:800;color:{color};line-height:1.25">{closing:,}개</div>'
                    f'<div style="font-size:12px;color:{color};margin-top:3px">{status}</div>',
                    unsafe_allow_html=True,
                )
            records.append({
                "inventory_date": as_of.isoformat(), "item_key": key,
                "opening_qty": int(opening), "received_qty": 0,
                "auto_ship_qty": auto[key], "actual_ship_qty": int(actual),
                "closing_qty": closing,
            })

    # 오늘 화면에 표시된 값(저장 전 수정사항 포함)을 한 번 눌러 즉시 복사.
    with copy_all_slot.container():
        render_oneclick_copy_button(
            make_message(as_of, records, compact=False),
            "📋 출고 + 재고 복사", "sikhye-copy-all", accent=True,
        )
    with copy_stock_slot.container():
        render_oneclick_copy_button(
            make_message(as_of, records, compact=True),
            "📋 재고만 복사", "sikhye-copy-stock",
        )
    with preview_slot.container():
        with st.popover("👁 미리보기", use_container_width=True):
            st.caption("출고 + 남은 재고")
            st.code(make_message(as_of, records, compact=False), language=None)
            st.divider()
            st.caption("칠판용 · 남은 재고만")
            st.code(make_message(as_of, records, compact=True), language=None)
            st.caption("※ 복사는 DB 저장과 별개입니다. 확정하려면 아래 최종 저장을 눌러주세요.")

    low_rows = [r for r in records if _is_low(r["item_key"], r["closing_qty"])]
    negative = any(r["closing_qty"] < 0 for r in records)
    if low_rows:
        st.warning("⚠️ 재고 부족 " + str(len(low_rows)) + "종: " + ", ".join(LABELS[r["item_key"]] for r in low_rows))
    else:
        st.success("✅ 모든 품목이 설정한 재고 기준 이상입니다.")
    if negative:
        st.error("출고량이 기존 재고보다 많은 품목이 있습니다. 수량을 확인해야 저장할 수 있습니다.")

    # 기존 스키마는 opening_qty + received_qty - actual_ship_qty = closing_qty를 사용하므로
    # received_qty는 0으로 저장하고, 생산량은 opening_qty에 직접 반영한다.
    note = st.text_input(
        "변경 메모 (선택)", key=_draft_key(as_of, "all", "note_v13"),
        placeholder="예: 생산 후 실물 재고로 수정 / 택배 외 8개 추가 출고",
    )
    if st.button("💾 오늘 재고 최종 저장", type="primary", use_container_width=True,
                 disabled=negative or bool(error)):
        try:
            expected_versions = {key: str(today_saved.get(key, {}).get("updated_at", "")) for key in KEYS}
            audit_ok = save_inventory(records, note, expected_versions=expected_versions)
            st.success("오늘 5개 품목의 실출고량과 남은 재고를 날짜별로 확정 저장했습니다. 같은 날 재저장하면 최신 값으로 수정됩니다.")
            if not audit_ok:
                st.warning("최종 재고는 저장됐지만 변경 이력 저장은 실패했습니다.")
            st.rerun()
        except Exception as exc:
            st.error(f"저장 실패: {type(exc).__name__} - {str(exc)[:150]}")

    st.divider()
    render_verified_inventory_history(history_rows, as_of)

    with st.expander("📈 최근 주간 출고 추이 (새 재고 저장 이력)", expanded=False):
        weeks = st.select_slider("조회 기간", options=[4, 8, 12, 26, 52], value=12,
                                 format_func=lambda n: f"최근 {n}주")
        weekly, avg = snapshots_weekly(history_rows, weeks=weeks, as_of=as_of)
        if weekly.empty:
            st.info("아직 저장된 재고 기록이 없습니다. 앞으로 실제 출고량을 저장하면 주간 기록이 쌓입니다.")
        else:
            pivot = weekly.pivot_table(index="주", columns="품목", values="출고수량", aggfunc="sum", fill_value=0)
            st.bar_chart(pivot)
            st.dataframe(weekly, use_container_width=True, hide_index=True)
            if not avg.empty:
                st.caption("월~금 5일 모두 기록한 주만 평균 계산에 포함")
                st.dataframe(avg.round(1), use_container_width=True, hide_index=True)
            else:
                st.caption("월~금 5일 모두 기록된 주가 없어 아직 평균을 표시하지 않습니다.")

    with st.expander("🕒 지난 재고 저장 내역", expanded=False):
        if not history_rows:
            st.info("저장 내역이 없습니다.")
        else:
            history = pd.DataFrame(history_rows)
            history["품목"] = history["item_key"].map(LABELS)
            history["기존재고"] = history["opening_qty"].fillna(0) + history["received_qty"].fillna(0)
            fields = ["inventory_date", "품목", "기존재고", "actual_ship_qty", "closing_qty", "updated_at"]
            st.dataframe(history[[x for x in fields if x in history.columns]].rename(columns={
                "inventory_date": "날짜", "actual_ship_qty": "오늘출고",
                "closing_qty": "남은재고", "updated_at": "저장시각",
            }), use_container_width=True, hide_index=True)


# =========================
# 앱 시작
# =========================
# =========================
# V27 · 내장 사용설명서 (기존 CRM 데이터 및 계산 로직 무변경)
# =========================
_GUIDE = {
    "🚀 빠른 시작": {
        "intro": "매일 가장 자주 쓰는 작업 순서예요. 처음 사용하는 사람은 여기부터 시작하세요.",
        "steps": [
            ("📂", "주문 업로드", "엑셀 업로드 탭에 쿠팡·네이버 주문 파일을 드래그해 넣습니다."),
            ("🔥", "오늘 주문 확인", "신규·재구매·VIP·주의 고객과 품목별 출고량을 확인합니다."),
            ("💾", "주문 DB 저장", "내용이 정확하면 엑셀 업로드 탭의 'DB에 저장하기'를 누릅니다."),
            ("📦", "재고 확정", "오늘 출고와 기존 재고를 확인·수정한 뒤 '오늘 재고 최종 저장'을 누릅니다."),
        ],
        "tips": ["주문 DB 저장과 재고 DB 저장은 서로 다른 버튼입니다.", "송장 작업은 '🚚 송장 업무'에서 별도로 진행하고 주문 DB에는 자동 저장되지 않습니다."],
    },
    "⬆️ 엑셀 업로드": {
        "intro": "쿠팡 DeliveryList 또는 네이버 스마트스토어 주문 엑셀을 분석하고 주문 DB에 반영하는 화면입니다.",
        "steps": [
            ("📑", "파일 준비", "쿠팡·네이버에서 주문 엑셀을 내려받습니다."),
            ("☁️", "한 번에 드래그", "큰 업로드 영역에 XLSX·XLS·CSV 파일 여러 개를 넣습니다."),
            ("🔎", "검증", "성공·실패 파일 수, 중복 가능 주문, 상품 수량을 확인합니다."),
            ("💾", "직접 저장", "확인 후 'DB에 저장하기'를 눌러 누적 주문 기록에 반영합니다."),
        ],
        "tips": ["업로드만으로 주문 DB에 저장되는 것은 아닙니다.", "같은 주문번호를 다시 저장하면 기존 기록이 갱신될 수 있으므로 날짜·주문내용을 확인하세요.", "쿠팡에서 주문번호가 같아도 상품행이 다를 수 있습니다. 중복 경고를 무시하지 마세요."],
    },
    "🚚 송장 업무": {
        "intro": "주문·송장 파일을 CJ 업무에 맞춰 변환하거나, 두 플랫폼의 송장번호를 자동으로 채웁니다.",
        "steps": [
            ("1️⃣", "네이버 1행 제거", "송장 업무 상단에 네이버 발주·발송관리 파일 1개를 올리면 비밀번호 1111 해제와 안내문 행 제거를 처리합니다."),
            ("📥", "4개 파일 넣기", "하단에 네이버 주문, 쿠팡 주문, 네이버 CJ 송장, 쿠팡 CJ 송장 엑셀 4개를 순서 상관없이 올립니다."),
            ("🛡️", "일치 검사", "이름·전화번호·주소·상품·수량이 일치하는지 확인합니다. 불일치 시 완성 파일을 만들지 않습니다."),
            ("📦", "한 번에 다운로드", "검사에 통과하면 버튼으로 ZIP 한 개를 받아 ★네이버 완성 / ★쿠팡 완성 두 파일을 꺼냅니다."),
        ],
        "tips": ["상단 '1행 제거' 기능과 하단 '4개 자동 매칭' 기능은 별개입니다.", "네이버 최종 파일은 발송처리 시트와 주문정보까지 보존한 암호 없는 .xls입니다. 네이버에 업로드한 다음 반드시 일괄 발송처리를 클릭하세요.", "ZIP 안의 송장번호를 최종 확인하세요. 파일은 고객·주문 DB에 자동 저장되지 않습니다."],
    },
    "🔥 오늘 주문 분석": {
        "intro": "오늘 올린 주문에서 신규·재구매·VIP·주의 고객과 품목별 출고를 확인합니다.",
        "steps": [
            ("⬆️", "주문 파일 업로드", "먼저 '엑셀 업로드' 탭에서 파일을 올립니다."),
            ("👥", "고객 구분", "신규·기존·VIP·주의·블랙 상태를 확인합니다."),
            ("📦", "품목별 수량", "출고 총합을 확인합니다. 묶음 구성은 실제 병 수량 기준인지 확인하세요."),
            ("🔍", "상세 확인", "고객별 주문 내역을 보고 이상이 있는 주문을 확인합니다."),
        ],
        "tips": ["'오늘 주문 분석'은 현재 업로드한 파일을 기반으로 합니다.", "업로드가 안 된 주문이나 송장에 누락된 출고는 별도로 확인해야 합니다."],
    },
    "📦 재고관리": {
        "intro": "5개 품목의 실제 출고와 남은 재고를 직접 확정해 저장하는 곳입니다.",
        "steps": [
            ("🧮", "출고 수량", "엑셀 자동 출고 수량을 확인하고 송장에 없는 출고를 포함해 '오늘 출고'를 수정합니다."),
            ("✍️", "기존 재고", "오늘 출고 전 재고를 입력합니다. 생산 후에는 증가한 총수량으로 수정하세요."),
            ("🚨", "부족 알림", "출고 후 재고가 품목별 기준 아래인지 바로 확인합니다."),
            ("💾", "최종 저장", "검토 후 '오늘 재고 최종 저장' 클릭. 상단 두 복사 버튼으로 카톡·칠판용 문구를 바로 복사합니다."),
        ],
        "tips": ["부족 기준: 식혜 1L 150 미만 / 500ml 100 이하 / 200ml 100 미만 / 단호박 1L 60 미만 / 단호박 500ml 30 미만.", "입력값을 바꾸거나 카톡에 복사하는 것만으로 DB에 저장되지는 않습니다.", "실출고를 날짜별로 매일 저장해야 향후 주간 생산 비교가 정확해집니다."],
    },
    "📊 대시보드": {
        "intro": "주문 성장률, 재구매, 판매 채널, 실제 출고 기반 생산 참고량을 분석합니다.",
        "steps": [
            ("📈", "판매·전년 비교", "이번 달 동기간, 지난달 전체, 올해 누적의 주문 증감과 월·요일 패턴을 확인합니다."),
            ("👥", "신규·재구매", "신규 고객 비중과 30·60·90일 재구매 전환율을 확인합니다."),
            ("🏭", "생산·재고", "매일 저장한 실제 출고 기록을 기반으로 품목별 생산 참고량과 소진 전망을 봅니다."),
            ("🛒", "채널·명절", "쿠팡과 네이버 비교, 설·추석 전 수요 변화를 살펴봅니다."),
        ],
        "tips": ["주문 건수와 실제 출고 병 수량은 다릅니다.", "기록이 부족하면 생산 예측은 '기록 부족'으로 표시될 수 있습니다.", "전년도 수치나 분석 결과를 보려면 주문 데이터가 DB에 누적돼 있어야 합니다."],
    },
    "👤 고객 CRM": {
        "intro": "저장된 전체 주문을 고객별로 묶어 구매횟수·최근 주문·재구매 현황을 살펴봅니다.",
        "steps": [
            ("🗃️", "주문 DB 기반", "업로드 화면에서 주문을 저장하면 고객 CRM에 누적 반영됩니다."),
            ("🔎", "검색·정렬", "고객 이름과 구매 횟수, 최근 주문을 기준으로 확인합니다."),
            ("♻️", "재구매 확인", "구매가 1회인지 2회 이상인지와 VIP 기준을 구분합니다."),
        ],
        "tips": ["고객 정보는 개인정보이므로 화면 공유와 엑셀 파일 전달 시 주의하세요.", "최근 업로드한 주문을 DB에 저장하지 않았다면 전체 고객 통계에 반영되지 않습니다."],
    },
    "🔎 고객 상세": {
        "intro": "특정 고객의 이전 주문과 메모를 자세히 확인합니다.",
        "steps": [
            ("🔍", "고객 찾기", "저장된 주문 DB에서 고객을 검색합니다."),
            ("📜", "주문 이력", "주문 날짜·횟수·상품 기록을 살펴봅니다."),
            ("📝", "메모 저장", "필요하면 고객 메모를 기록하고 저장 버튼으로 확정합니다."),
        ],
        "tips": ["동명이인일 수 있으므로 고객 식별정보를 함께 확인하세요.", "메모에는 불필요한 민감정보를 넣지 않는 것이 좋습니다."],
    },
    "⭐ VIP": {
        "intro": "반복해서 주문하는 단골 고객을 확인하는 화면입니다.",
        "steps": [
            ("⭐", "VIP 목록", "기존 주문기록 기준 구매 5회 이상 고객을 확인합니다."),
            ("🛍️", "주문 패턴", "고객의 주문 빈도와 최근 구매를 확인합니다."),
            ("🔎", "상세 조회", "궁금한 고객은 '고객 상세'에서 주문 이력을 찾아봅니다."),
        ],
        "tips": ["VIP 등급은 저장된 주문 데이터가 충분해야 정확합니다."],
    },
    "🤖 AI CRM 비서": {
        "intro": "CRM 통계를 바탕으로 판매 흐름을 해석하거나 질문하는 보조 분석 기능입니다.",
        "steps": [
            ("✨", "자동 분석", "AI CRM 비서에서 자동 분석 실행 버튼을 누릅니다."),
            ("💬", "질문하기", "판매 추세나 재구매 개선 방향 등을 질문합니다."),
            ("✅", "검토하기", "AI의 제안은 참고용입니다. 실제 주문·출고량은 대시보드와 재고 기록을 우선 확인합니다."),
        ],
        "tips": ["AI 키·모델 설정 또는 이용량 제한에 따라 작동하지 않을 수 있습니다.", "개인정보가 포함된 주문 원본을 AI에 별도 붙여넣지 마세요."],
    },
    "🚫 블랙리스트": {
        "intro": "배송·고객 응대 시 확인해야 할 주의 고객과 사유를 관리합니다.",
        "steps": [
            ("🚦", "상태 확인", "주의·블랙 상태를 확인합니다."),
            ("📝", "사유 기록", "필요한 고객을 선택해 정확한 사유를 기록합니다."),
            ("💾", "저장·삭제", "기록 변경 후 결과를 확인합니다."),
        ],
        "tips": ["고객 오인 방지를 위해 같은 이름의 고객인지 확인하세요.", "잘못된 기록을 삭제할 때는 대상 고객을 다시 확인하세요."],
    },
    "📦 주문 DB · ⬇️ 다운로드": {
        "intro": "저장된 주문을 조회하고, 보고·백업용 CRM 엑셀을 내려받습니다.",
        "steps": [
            ("🗄️", "주문 DB", "지금까지 저장된 누적 주문을 확인합니다."),
            ("📊", "자료 선택", "다운로드 탭에서 요약·고객·VIP·재구매·주문 기록을 준비합니다."),
            ("📥", "엑셀 생성", "'CRM 전체 엑셀 파일 만들기'를 누른 뒤 준비된 파일을 다운로드합니다."),
        ],
        "tips": ["고객 개인정보가 담긴 엑셀 파일은 공유 폴더에 무분별하게 보관하지 마세요.", "다운로드한 엑셀의 수정 내용은 프로그램 DB에 자동 반영되지 않습니다."],
    },
    "🛠️ 문제 해결": {
        "intro": "잘 안 될 때 가장 먼저 확인할 내용을 모았습니다.",
        "steps": [
            ("❌", "엑셀 업로드 실패", "지원 확장자·파일 손상·암호 여부를 확인하고 파일별 처리 결과를 펼쳐보세요."),
            ("🔄", "저장 후 숫자 그대로", "올바른 저장 버튼을 눌렀는지 확인하고 DB 기반 화면을 다시 열어보세요."),
            ("🚚", "송장 매칭 실패", "주문 2개와 CJ 송장 2개가 맞는 날짜인지, 수취인·연락처·주소·상품·수량이 일치하는지 확인하세요."),
            ("☁️", "DB·AI 연결 오류", "네트워크 상태와 오류 문구를 확인하고, 관리자에게 오류 내용을 전달하세요. API 키는 공유하지 마세요."),
        ],
        "tips": ["고객정보가 있는 원본 엑셀이나 API 비밀키를 공개 채팅방·스크린샷에 올리지 마세요.", "송장 매칭 오류가 나면 임의로 행 순서만 맞춰서 저장하지 마세요."],
    },
}


def _guide_steps_html(steps):
    """사용설명서 그림형 흐름도: 외부 파일 없이 Streamlit에서 자체 렌더링."""
    import html as _html
    cells = []
    for i, (symbol, title, desc) in enumerate(steps, 1):
        cells.append(
            '<div class="guide-step">'
            f'<div class="guide-illustration" aria-hidden="true">{_html.escape(symbol)}</div>'
            f'<div class="guide-step-num">STEP {i:02d}</div>'
            f'<div class="guide-step-name">{_html.escape(title)}</div>'
            f'<div class="guide-step-desc">{_html.escape(desc)}</div>'
            '</div>'
        )
    return '<div class="guide-steps-grid">' + ''.join(cells) + '</div>'


@st.dialog("📘 식혜명가 CRM 사용설명서", width="large")
def show_crm_guide():
    st.markdown("""
    <style>
    .guide-steps-grid {display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:11px;margin:12px 0 16px;}
    .guide-step {border:1px solid #e2e8f0;border-radius:14px;background:#f8fafc;padding:14px;min-height:164px;}
    .guide-illustration {font-size:34px;line-height:1.25;margin-bottom:8px;}
    .guide-step-num {font-size:11px;font-weight:800;color:#2563eb;letter-spacing:.05em;}
    .guide-step-name {font-size:17px;font-weight:800;color:#0f172a;margin:4px 0 6px;}
    .guide-step-desc {font-size:14px;color:#475569;line-height:1.6;}
    @media(max-width:600px){.guide-steps-grid{grid-template-columns:1fr}.guide-step{min-height:auto}}
    </style>
    """, unsafe_allow_html=True)
    st.caption("원하는 메뉴를 선택하면 사용 순서가 그림과 함께 나옵니다. 별도 파일이나 외부 링크는 필요 없어요.")
    guide_names = list(_GUIDE)
    selected = st.selectbox(
        "📚 설명서 목차 · 알고 싶은 기능 선택",
        guide_names,
        key="_crm_manual_selected_v27",
    )
    info = _GUIDE[selected]
    st.markdown(f"### {selected}")
    st.write(info["intro"])
    st.markdown(_guide_steps_html(info["steps"]), unsafe_allow_html=True)
    st.markdown("**💡 꼭 알아둘 점**")
    for tip in info["tips"]:
        st.markdown("- " + tip)
    st.divider()
    st.caption("사용설명서는 실제 CRM 탭 이름과 저장 방식을 기준으로 작성했습니다. 사용설명서를 열어도 데이터가 변경되거나 저장되지 않습니다.")


header_title, header_help = st.columns([5.3, 1.2], vertical_alignment="center")
with header_title:
    st.title("📦 식혜명가 저장형 고객 CRM")
with header_help:
    if st.button("📘 사용설명서", key="crm_open_manual_v27", use_container_width=True,
                 help="기능별 사용법과 그림 설명 보기"):
        show_crm_guide()
st.caption("쿠팡·네이버 주문부터 재고·고객·전년 실적까지 한 곳에서 확인합니다.")

if sb is None:
    st.markdown("""
<div class="warn">
<b>Supabase 연결이 아직 안 됐습니다.</b><br>
Streamlit Secrets에 SUPABASE_URL, SUPABASE_KEY를 넣어야 저장형 CRM으로 작동합니다.<br>
설정 전에는 앱 화면만 확인할 수 있습니다.
</div>
""", unsafe_allow_html=True)
    st.stop()


tab_upload, tab_shipping, tab_today, tab_inventory, tab_dashboard, tab_customers, tab_detail, tab_vip, tab_ai, tab_black, tab_orders, tab_download = st.tabs([
    "⬆️ 엑셀 업로드", "🚚 송장 업무", "🔥 오늘 주문 분석", "📦 재고관리", "📊 대시보드", "👤 고객 CRM", "🔎 고객 상세", "⭐ VIP", "🤖 AI CRM 비서", "🚫 블랙리스트", "📦 주문 DB", "⬇️ 다운로드"
], on_change="rerun", key="crm_main_tabs")

if "today_order_preview" not in st.session_state:
    st.session_state["today_order_preview"] = pd.DataFrame()
if "today_raw_orders" not in st.session_state:
    st.session_state["today_raw_orders"] = pd.DataFrame()
if "today_customer_analysis" not in st.session_state:
    st.session_state["today_customer_analysis"] = pd.DataFrame()
if "today_detail_analysis" not in st.session_state:
    st.session_state["today_detail_analysis"] = pd.DataFrame()

# 업로드 위젯은 항상 렌더링해 탭을 옮겼다가 돌아와도 파일 선택이 유지되게 한다.
# 실제 주문 파싱·분석은 업로드 탭을 열었을 때에만 수행한다.
with tab_upload:
    st.markdown("""
    <div class="upload-guide" style="padding:14px 16px;margin-bottom:9px;border:1px solid #dbeafe;background:#f0f7ff;">
        <div class="upload-title" style="font-size:21px;">📂 쿠팡 · 네이버 주문 엑셀</div>
        <div class="upload-sub" style="font-size:15px;">아래 커다란 영역 어디든 끌어놓거나 <b>Upload 버튼</b>을 클릭하세요. 여러 파일을 한꺼번에 넣어도 됩니다.</div>
    </div>
    """, unsafe_allow_html=True)

    uploaded_files = st.file_uploader(
        "➕ 엑셀 파일 추가 업로드",
        type=["xlsx", "xls", "csv"],
        accept_multiple_files=True,
        help="쿠팡 DeliveryList, 네이버 스마트스토어 엑셀을 여러 개 동시에 올릴 수 있습니다.",
        key=f"crm_excel_uploader_v21_{st.session_state.get('_upload_widget_version_v21', 0)}",
    )

if tab_upload.open:
    with tab_upload:
        if uploaded_files:
            # Streamlit은 위젯 변경 시 스크립트를 재실행한다. 동일 업로드 파일은 분석결과를 재사용한다.
            import hashlib
            blobs = tuple((f.name, f.getvalue()) for f in uploaded_files)
            signature = tuple((name, len(content), hashlib.sha256(content).hexdigest())
                              for name, content in blobs)
            with st.spinner("엑셀 주문 내역을 분석하는 중입니다. 같은 파일은 다시 계산하지 않습니다..."):
                raw_orders, order_preview, read_logs, errors = _prepare_uploaded_files(blobs)
            log_df = pd.DataFrame(read_logs)
            success_count = int((log_df["상태"] == "성공").sum()) if not log_df.empty else 0
            fail_count = int((log_df["상태"] == "실패").sum()) if not log_df.empty else 0
            st.markdown(f'<div class="summary-box">✅ 읽기 완료 · 성공 <b>{success_count:,}</b>개 / 실패 <b>{fail_count:,}</b>개</div>',
                        unsafe_allow_html=True)
            with st.expander("파일별 처리 결과"):
                if not log_df.empty:
                    st.dataframe(log_df, use_container_width=True, hide_index=True)
            for name, msg in errors:
                st.error(f"{name}: {msg}")
            if not raw_orders.empty and len(raw_orders["source_file"].unique()) > 1:
                # 동일 주문이 서로 다른 쿠팡 엑셀에 나타날 경우 서로 다른 상품행일 수도 있다.
                # 자동 삭제하지 않고 중복 출고 가능성을 안내해 사람이 확인하도록 한다.
                cp = raw_orders[raw_orders["channel"].eq("쿠팡")]
                if not cp.empty:
                    overlap = cp.groupby("order_key")["source_file"].nunique()
                    count = int(overlap.gt(1).sum())
                    if count:
                        st.warning(f"서로 다른 쿠팡 파일에 동일 주문번호 {count:,}건이 있습니다. "
                                   "수량이 중복 계산될 수 있으니 오늘 출고 수량을 확인해 주세요.")

            if not raw_orders.empty:
                # 고객 신규/재구매 판정은 파일 변경 또는 주문 DB 갱신 시에만 다시 실행한다.
                orders_db_current = fetch_orders_as_order_level()
                db_signature = (len(orders_db_current),
                                str(orders_db_current["updated_at"].max()) if "updated_at" in orders_db_current.columns and not orders_db_current.empty else "")
                analysis_signature = (signature, db_signature,
                                      st.session_state.get("_orders_revision_v21", 0),
                                      st.session_state.get("_crm_customer_flags_revision_v21", 0))
                if st.session_state.get("_upload_analysis_signature_v21") != analysis_signature:
                    customer_df_current = make_customer_df(orders_db_current) if not orders_db_current.empty else pd.DataFrame()
                    black_df_current = fetch_blacklist()
                    today_customer_analysis, today_detail_analysis = analyze_today_orders(
                        order_preview, orders_db_current, customer_df_current, black_df_current)
                    st.session_state["today_raw_orders"] = raw_orders
                    st.session_state["today_order_preview"] = order_preview
                    st.session_state["today_customer_analysis"] = today_customer_analysis
                    st.session_state["today_detail_analysis"] = today_detail_analysis
                    st.session_state["_upload_analysis_signature_v21"] = analysis_signature
                today_customer_analysis = st.session_state["today_customer_analysis"]
                today_detail_analysis = st.session_state["today_detail_analysis"]
                st.subheader("오늘 업로드 요약")
                c1, c2, c3, c4, c5 = st.columns(5)
                c1.metric("원본 상품행", f"{len(raw_orders):,}행")
                c2.metric("실제 주문", f"{len(order_preview):,}건")
                c3.metric("예상 고객", f"{order_preview['analysis_customer_key'].nunique():,}명")
                c4.metric("오늘 재구매", f"{int((today_customer_analysis['기존주문횟수'] >= 1).sum()) if not today_customer_analysis.empty else 0:,}명")
                c5.metric("주의/블랙", f"{int((today_customer_analysis['블랙상태'].isin(['주의','블랙'])).sum()) if not today_customer_analysis.empty else 0:,}명")

                st.success("분석 완료. 위쪽의 `🔥 오늘 주문 분석` 탭에서 재구매/VIP/블랙 고객을 색상으로 확인하세요.")

                with st.expander("오늘 주문 고객 요약 바로 보기", expanded=True):
                    if not today_customer_analysis.empty:
                        st.dataframe(style_today_customer_table(today_customer_analysis), use_container_width=True, hide_index=True)
                    else:
                        st.info("오늘 주문 분석 데이터가 없습니다.")

                with st.expander("저장 전 주문 단위 상세 보기", expanded=False):
                    st.dataframe(order_preview.sort_values("order_date", ascending=False), use_container_width=True, hide_index=True)

                col_save, col_clear = st.columns([1, 1])
                with col_save:
                    if st.button("DB에 저장하기", type="primary", use_container_width=True):
                        saved = upsert_orders(order_preview)
                        st.success(f"{saved:,}건 저장 완료. 대량 자료는 1000건씩 시도하고, 실패 시 자동으로 더 작게 나눠 저장했습니다. 기존 주문번호는 자동으로 중복 제거/갱신되었습니다.")
                        st.info("저장 후 대시보드/고객 CRM 탭을 새로고침하면 누적 DB 기준으로 반영됩니다.")
                with col_clear:
                    if st.button("화면 분석 결과 초기화", use_container_width=True):
                        st.session_state["today_raw_orders"] = pd.DataFrame()
                        st.session_state["today_order_preview"] = pd.DataFrame()
                        st.session_state["today_customer_analysis"] = pd.DataFrame()
                        st.session_state["today_detail_analysis"] = pd.DataFrame()
                        st.session_state.pop("_upload_analysis_signature_v21", None)
                        st.session_state["_upload_widget_version_v21"] = st.session_state.get("_upload_widget_version_v21", 0) + 1
                        st.rerun()
        elif not st.session_state.get("today_raw_orders", pd.DataFrame()).empty:
            st.info("✅ 이전 업로드 분석이 이번 로그인 세션에 남아 있어요. 오늘 주문 분석·재고관리 탭에서 계속 사용할 수 있습니다. 새 파일을 올리면 갱신됩니다.")

# 송장 업무는 일반 주문/고객/재고 DB와 완전히 별도. 파일은 저장하지 않고 브라우저 다운로드만 제공.
# 위: 네이버 안내문 1행 제거 / 아래: 네이버+쿠팡+CJ 4개 파일 동시 매칭.
with tab_shipping:
    st.subheader("🚚 송장 업무 자동화")
    st.caption("엑셀을 직접 수정하지 않고, 파일 업로드와 다운로드로 송장 업무를 끝냅니다. DB에는 저장되지 않습니다.")

    st.markdown("### ① 네이버 엑셀 첫 줄 제거")
    st.caption("CJ에 주문 엑셀을 등록하기 전, 스마트스토어 원본의 안내문 첫 행만 제거합니다. 파일 암호 1111 자동 처리.")
    cj_uploaded = st.file_uploader(
        "네이버 발주·발송관리 파일 1개 드래그",
        type=["xlsx"],
        key="ship_clean_naver_v26",
        help="네이버 원본 엑셀만 넣어주세요. 이미 첫 행을 제거한 파일은 변환하지 않습니다.",
    )
    if tab_shipping.open and cj_uploaded is not None:
        try:
            with st.spinner("암호 해제 · 안내문 1행 삭제 · 데이터 검증 중..."):
                cleaned_bytes, clean_count, clean_cols = convert_naver_cj_excel(cj_uploaded.getvalue())
            st.success(f"변환 완료: 주문 {clean_count:,}행 · {clean_cols:,}개 열 유지")
            st.download_button(
                "💾 네이버 1행 제거 파일 저장",
                data=cleaned_bytes,
                file_name=f"변완 네이버 송장 ({datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d')}).xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary", use_container_width=True,
                key="ship_clean_download_v26",
            )
        except Exception as exc:
            st.error(f"변환하지 않았습니다: {exc}")

    st.divider()
    st.markdown("### ② 네이버 + 쿠팡 송장번호 자동 매칭")
    st.caption("파일 4개를 **순서 상관없이 한 번에 드래그**하세요. 네이버 최종 파일은 암호 없는 발송처리 시트의 .xls, 쿠팡은 .xlsx로 저장됩니다.")
    st.info("📂 필요한 파일: 네이버 주문 1개 · 쿠팡 주문 1개 · 네이버 CJ 송장 1개 · 쿠팡 CJ 송장 1개")
    joined_uploads = st.file_uploader(
        "📥 주문 및 CJ 송장 엑셀 4개 동시 업로드",
        type=["xlsx"],
        accept_multiple_files=True,
        key="ship_four_files_v26",
        help="네이버 파일 암호 1111 자동 해제. 송장번호와 수취인 주소·전화·상품·수량을 검증한 뒤 결합합니다.",
    )
    if tab_shipping.open and joined_uploads:
        if len(joined_uploads) != 4:
            st.warning(f"현재 {len(joined_uploads)}개 선택됨 · 주문 2개와 CJ 송장 2개, 총 4개를 넣어주세요.")
        else:
            try:
                import hashlib
                data_files = [(item.name, item.getvalue()) for item in joined_uploads]
                signature = (datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y%m%d"),
                             tuple(sorted((name, hashlib.sha256(blob).hexdigest()) for name, blob in data_files)))
                if st.session_state.get("_ship_v26_signature") != signature:
                    with st.spinner("파일 자동 판별 → 이름·연락처·주소·상품·수량 확인 → 송장번호 입력 → 최종 엑셀 재검사 중..."):
                        datecode = datetime.now(ZoneInfo('Asia/Seoul')).strftime('%Y%m%d')
                        zip_bytes, summary, problems = make_shipment_zip(data_files, datecode)
                    st.session_state["_ship_v26_signature"] = signature
                    st.session_state["_ship_v26_result"] = (zip_bytes, summary, problems, datecode)
                else:
                    zip_bytes, summary, problems, datecode = st.session_state["_ship_v26_result"]

                for platform in ("네이버", "쿠팡"):
                    stats = summary[platform]
                    st.markdown(
                        f"**{platform}**: 주문 **{stats['orders']:,}행** · CJ 송장 **{stats['carriers']:,}행** "
                        f"· 매칭 **{stats['matched']:,}행** · 실제 송장번호 **{stats['unique_tracks']:,}개**"
                    )
                if problems:
                    st.error("🚨 일치하지 않는 행이 있어 최종 파일 생성을 중단했습니다. 파일 4개를 확인해 주세요.")
                    with st.expander("어느 행을 확인해야 하나요?", expanded=True):
                        for problem in problems[:35]:
                            st.write("- " + problem)
                        if len(problems) > 35:
                            st.caption(f"그 외 {len(problems)-35:,}건은 파일 재확인이 필요합니다.")
                else:
                    st.success("✅ 두 플랫폼 모두 검증 완료. 네이버는 주문원본 열을 보존한 발송처리 .xls(암호 없음), 쿠팡은 원본 형식 .xlsx로 만들었습니다.")
                    st.download_button(
                        "📦 ★네이버 + ★쿠팡 완성 파일 2개 한 번에 저장",
                        data=zip_bytes,
                        file_name=f"송장완성_네이버+쿠팡_{datecode}.zip",
                        mime="application/zip",
                        type="primary", use_container_width=True,
                        key="ship_zip_download_v26",
                    )
                    st.caption(f"한 번 클릭하면 ZIP이 다운로드됩니다. 압축을 풀면 **★네이버 완성({datecode}).xls**, **★쿠팡 완성({datecode}).xlsx** 두 파일이 나옵니다.")
            except Exception as exc:
                st.session_state.pop("_ship_v26_signature", None)
                st.session_state.pop("_ship_v26_result", None)
                st.error(f"송장 매칭을 중단했습니다: {exc}")
                st.caption("잘못된 송장번호를 넣지 않기 위해 확인되지 않은 파일은 다운로드하지 않습니다.")

# 업로드 탭을 보고 있을 때는 다른 탭의 1만 건 이상 DB/차트 분석을 실행하지 않는다.
# 파일을 올린 경우 신규·재구매 분석에 필요한 주문 DB만 위쪽에서 별도로 확인한다.
if tab_upload.open or tab_shipping.open:
    st.stop()

# DB 데이터 로드: 재고관리는 실제 출고 스냅샷만 필요하므로 주문 1만여 건 전체 조회 생략.
# 다른 모든 탭은 기존 주문 DB를 읽어 기존 동작을 유지한다.
orders_db = pd.DataFrame() if tab_inventory.open else fetch_orders_as_order_level()
black_df = (fetch_blacklist() if any(t.open for t in
            (tab_dashboard, tab_customers, tab_detail, tab_vip, tab_ai, tab_black, tab_download))
            else pd.DataFrame())

if orders_db.empty:
    if tab_inventory.open:
        with tab_inventory:
            render_inventory(
                raw_orders=st.session_state.get("today_raw_orders", pd.DataFrame()),
                orders_db=orders_db,
                shipping_summary_func=make_shipping_summary,
                infer_shipping_item=infer_shipping_item,
            )
    if tab_dashboard.open:
        with tab_dashboard:
            st.info("아직 저장된 주문 DB가 없습니다. 먼저 엑셀을 업로드하고 DB에 저장하세요.")
    if tab_today.open or tab_customers.open:
        st.info("저장된 주문 자료가 없습니다. 엑셀 업로드에서 저장을 확인해 주세요.")
    st.stop()

needs_customer_profile = any(t.open for t in (tab_dashboard, tab_customers, tab_detail, tab_vip, tab_ai, tab_black, tab_download))
if needs_customer_profile:
    customer_df = make_customer_df(orders_db)

    if not black_df.empty:
        customer_df = customer_df.merge(
            black_df[["customer_key", "status", "reason", "memo", "incident_date", "result"]],
            left_on="analysis_customer_key",
            right_on="customer_key",
            how="left"
        )
    else:
        customer_df["status"] = ""
        customer_df["reason"] = ""
        customer_df["memo"] = ""
        customer_df["incident_date"] = ""
        customer_df["result"] = ""

    customer_df["black_status"] = customer_df["status"].fillna("정상").replace("", "정상")
    customer_df["총주문금액"] = customer_df["total_amount"]
    customer_df["건당평균금액"] = customer_df["avg_order_amount"]

else:
    customer_df = pd.DataFrame()

if tab_inventory.open:
    with tab_inventory:
        render_inventory(
            raw_orders=st.session_state.get("today_raw_orders", pd.DataFrame()),
            orders_db=orders_db,
            shipping_summary_func=make_shipping_summary,
            infer_shipping_item=infer_shipping_item,
        )


if tab_today.open:
    with tab_today:
        st.subheader("🔥 오늘 주문 분석")
        st.caption("엑셀 업로드 후 DB 저장 전에도 기존 DB와 비교해서 오늘 주문 고객이 신규인지 재구매인지 바로 확인합니다.")

        today_customer_analysis = st.session_state.get("today_customer_analysis", pd.DataFrame())
        today_detail_analysis = st.session_state.get("today_detail_analysis", pd.DataFrame())

        if today_customer_analysis is None or today_customer_analysis.empty:
            st.info("먼저 `엑셀 업로드` 탭에서 오늘 쿠팡/네이버 엑셀을 업로드하세요.")
        else:
            total_today_customers = len(today_customer_analysis)
            reorder_count = int((today_customer_analysis["기존주문횟수"] >= 1).sum())
            new_count = int((today_customer_analysis["기존주문횟수"] == 0).sum())
            vip_reorder_count = int(today_customer_analysis["오늘고객상태"].astype(str).str.contains("VIP").sum())
            black_today_count = int(today_customer_analysis["블랙상태"].isin(["주의", "블랙"]).sum())

            a, b, c, d, e = st.columns(5)
            a.metric("오늘 고객", f"{total_today_customers:,}명")
            b.metric("오늘 신규", f"{new_count:,}명")
            c.metric("오늘 재구매", f"{reorder_count:,}명")
            d.metric("VIP 재주문", f"{vip_reorder_count:,}명")
            e.metric("주의/블랙", f"{black_today_count:,}명")

            st.markdown("""
            - ⚪ 신규: 기존 DB에 없던 고객
            - 🟢 재주문: 기존 주문 이력이 있는 고객
            - 🟡 VIP 재주문: 기존 5회 이상 구매 고객
            - 🔵 복귀 재주문: 90일 이상 미구매 후 재주문
            - 🔴 주의/블랙: 블랙리스트 또는 주의 고객
            """)

            st.subheader("📦 오늘 출고 총합")
            today_raw_orders = st.session_state.get("today_raw_orders", pd.DataFrame())
            shipping_summary = make_shipping_summary(today_raw_orders)

            if shipping_summary.empty:
                st.info("오늘 업로드된 주문이 없습니다.")
            else:
                st.dataframe(shipping_summary, use_container_width=True, hide_index=True)

            with st.expander("주별 출고량 / 품목별 주간 평균", expanded=False):
                weekly_table, weekly_avg = make_weekly_shipping_table_from_orders(today_raw_orders)

                st.markdown("**업로드 파일 기준 주별 출고량**")
                if weekly_table.empty:
                    st.info("주별 출고량 데이터가 없습니다.")
                else:
                    st.dataframe(weekly_table, use_container_width=True, hide_index=True)

                st.markdown("**품목별 주간 평균 출고량**")
                if weekly_avg.empty:
                    st.info("품목별 주간 평균 데이터가 없습니다.")
                else:
                    st.dataframe(weekly_avg, use_container_width=True, hide_index=True)

            with st.expander("저장 DB 기준 최근 N주 평균 출고량 (기존 주문 단위 추정)", expanded=False):
                st.warning("기존 주문 DB는 원본 상품행이 없어 혼합 주문의 실제 병 수량을 정확히 복원할 수 없습니다. 생산량 판단에는 새 재고관리 탭의 기록을 사용하세요.")
                avg_weeks = st.slider("분석 기간", min_value=2, max_value=26, value=8, step=1)
                weekly_avg_db = make_weekly_shipping_average(orders_db, weeks=avg_weeks)
                if weekly_avg_db.empty:
                    st.info("주간 평균 출고량 데이터가 없습니다.")
                else:
                    st.dataframe(weekly_avg_db, use_container_width=True, hide_index=True)

            filter_status = st.multiselect(
                "상태 필터",
                options=sorted(today_customer_analysis["오늘고객상태"].unique()),
                default=sorted(today_customer_analysis["오늘고객상태"].unique()),
            )

            view_today = today_customer_analysis[today_customer_analysis["오늘고객상태"].isin(filter_status)]
            st.dataframe(style_today_customer_table(view_today), use_container_width=True, hide_index=True, height=420)

            st.subheader("고객별 역대 주문내역 보기")
            selected_customer = st.selectbox(
                "고객 선택",
                options=view_today["고객키"].tolist(),
                format_func=lambda k: view_today.loc[view_today["고객키"] == k, "오늘고객명"].iloc[0] + " | " + str(view_today.loc[view_today["고객키"] == k, "오늘고객상태"].iloc[0])
            )

            history = orders_db[orders_db["analysis_customer_key"] == selected_customer].sort_values("order_date", ascending=False)
            st.write("기존 DB 저장 주문내역")
            if history.empty:
                st.info("기존 주문내역 없음. 이번 주문이 첫 주문입니다.")
            else:
                st.dataframe(history, use_container_width=True, hide_index=True)

            if today_detail_analysis is not None and not today_detail_analysis.empty:
                st.write("이번 업로드 주문 상세")
                today_detail_customer = today_detail_analysis[today_detail_analysis["analysis_customer_key"] == selected_customer]
                st.dataframe(today_detail_customer, use_container_width=True, hide_index=True)




# =========================
# V22 판매/고객/생산 경영 대시보드
# - 기존 DB에 쓰기 작업 없음
# - 주문 건수와 실제 병 수량을 명확하게 구분
# - 전년 비교는 같은 날짜까지, 미완료 기간은 별도 표기
# =========================
# =========================
# V22 경영 대시보드: 기존 DB를 읽기만 하는 순수 계산 함수
# =========================
@st.cache_data(ttl=180, max_entries=4, show_spinner=False)
def _v20_orders(orders):
    """주문 시각은 KST 기준 날짜로만 집계; 주문건수와 병 수량을 혼동하지 않음."""
    if orders is None or orders.empty or "order_date" not in orders.columns:
        return pd.DataFrame(columns=["local_day", "month", "weekday", "channel", "analysis_customer_key"])
    fields = [c for c in ["order_date", "channel", "analysis_customer_key", "receiver_name", "buyer_name"] if c in orders.columns]
    d = orders[fields].copy()
    parsed = pd.to_datetime(d["order_date"], errors="coerce", utc=True)
    d = d.loc[parsed.notna()].copy()
    parsed = parsed.loc[parsed.notna()]
    d["local_day"] = parsed.dt.tz_convert(KST).dt.date
    d["month"] = parsed.dt.tz_convert(KST).dt.strftime("%Y-%m")
    d["weekday"] = parsed.dt.tz_convert(KST).dt.weekday
    d["channel"] = d.get("channel", pd.Series("기타", index=d.index)).fillna("기타").astype(str).map(
        lambda x: "쿠팡" if "쿠팡" in x else ("네이버" if "네이버" in x else "기타"))
    d["analysis_customer_key"] = d.get("analysis_customer_key", pd.Series("", index=d.index)).fillna("").astype(str)
    return d


def _v20_period_count(sales, start_day, end_day):
    return int(sales["local_day"].between(start_day, end_day).sum()) if not sales.empty else 0


def _v20_compare(sales, days, today):
    # 오늘은 아직 진행 중이므로 어제까지 완료된 날짜만 사용
    end = today - timedelta(days=1)
    cur = _v20_period_count(sales, end - timedelta(days=days-1), end)
    prev = _v20_period_count(sales, end - timedelta(days=2*days-1), end - timedelta(days=days))
    pct = (cur - prev) / prev * 100 if prev > 0 else None
    return cur, prev, pct


def _v20_complete_weeks(inventory_rows, today, lookback=4):
    """오늘 포함 주는 제외; 5품목×월~금 모두 최종 저장된 주만 평균에 사용."""
    calendar = _recorded_inventory_calendar(inventory_rows)
    monday = today-timedelta(days=today.weekday())
    completed = []
    for i in range(lookback, 0, -1):
        info = _week_snapshot(calendar, monday-timedelta(weeks=i))
        if info["complete_days"] == 5:
            completed.append(info)
    return completed


def _v20_latest_stock(inventory_rows, today):
    found = {}
    for row in inventory_rows:
        try:
            day = date.fromisoformat(str(row.get("inventory_date", ""))[:10])
            key = row.get("item_key")
            if key not in KEYS or day > today:
                continue
            if key not in found or day > found[key][0]:
                found[key] = (day, _safe_int(row.get("closing_qty")))
        except (ValueError, TypeError):
            continue
    return found


def _v20_production(inventory_rows, today, min_weeks=2):
    completed = _v20_complete_weeks(inventory_rows, today, 4)
    stock = _v20_latest_stock(inventory_rows, today)
    report = []
    for key, label, threshold, inclusive in ITEMS:
        weeks = [w["items"][key]["total"] for w in completed]
        weekly_mean = sum(weeks)/len(weeks) if weeks else None
        record = stock.get(key)
        qty = record[1] if record else None
        freshness = (today-record[0]).days if record else None
        usable = len(weeks) >= min_weeks and qty is not None and freshness is not None and freshness <= 7
        # 필요량 = 다음 주 평일 예상 출고량 + 안전재고 기준 - 최신 재고
        safe_target = threshold + (1 if inclusive else 0)
        recommended = max(0, int(math.ceil(weekly_mean + safe_target - qty))) if usable else None
        expected_days = (qty / (weekly_mean/5)) if usable and weekly_mean and weekly_mean>0 else None
        report.append({"품목": label,
                       "최신 재고": qty if qty is not None else "—",
                       "재고 확인일": record[0].isoformat() if record else "—",
                       "최근 4주 확보": f"{len(weeks)}/4주",
                       "주간 예상 출고": round(weekly_mean) if len(weeks)>=min_weeks else "—",
                       "소진 예상(영업일)": round(expected_days,1) if expected_days is not None else "—",
                       "생산 참고량": recommended if recommended is not None else "—",
                       "기준": ("참고치" if usable else "기록 부족/재고 갱신 필요")})
    return pd.DataFrame(report), len(completed)


def _v20_holiday(sales, event_current, event_previous, today):
    """명절 날짜를 기준으로 D-21~D-1 세 구간을 비교. 주문건수로만 표시."""
    rows = []
    for first in ("3주 전", "2주 전", "1주 전"):
        index = {"3주 전":(21,15), "2주 전":(14,8), "1주 전":(7,1)}[first]
        prev_start = event_previous-timedelta(days=index[0])
        prev_end = event_previous-timedelta(days=index[1])
        cur_start = event_current-timedelta(days=index[0])
        cur_end = event_current-timedelta(days=index[1])
        old = _v20_period_count(sales, prev_start, prev_end)
        new = _v20_period_count(sales, cur_start, min(cur_end,today-timedelta(days=1))) if cur_start < today else 0
        complete = cur_end < today
        rows.append({"기간": first, f"{event_previous.year} 주문": old,
                     f"{event_current.year} 주문": new if complete else None,
                     "비교": (f"{(new-old)/old:+.1%}" if complete and old else "비교 불가")})
    return pd.DataFrame(rows)


def _v20_big_chart(chart, height=320, key=None, selectable=False):
    chart = chart.properties(height=height).configure_axis(
        labelFontSize=16, titleFontSize=17, labelPadding=10, titlePadding=12,
        gridColor="#e8eef5").configure_legend(labelFontSize=16, titleFontSize=16).configure_view(strokeWidth=0)
    kwargs={"use_container_width":True}
    if selectable:
        kwargs.update({"key":key,"on_select":"rerun"})
    return st.altair_chart(chart, **kwargs)


def _v20_delta(pct):
    return f"{pct:+.1f}%" if pct is not None else "이전 기간 0건"



def _v22_previous_year(day):
    """2월 29일은 전년의 2월 28일과 비교."""
    try:
        return day.replace(year=day.year - 1)
    except ValueError:
        return date(day.year - 1, 2, 28)


def _v22_pct(cur, prev):
    return round((cur - prev) / prev * 100, 1) if prev else None


def _v22_pct_label(pct):
    return f"{pct:+.1f}%" if pct is not None else "비교 불가"


def _v22_yoy(sales, start, end):
    """올해와 전년의 동일 월일 범위; 주기적 판매/캠페인 비용 데이터가 아님."""
    if start > end:
        return {"now": None, "before": None, "pct": None, "range": None}
    old_start, old_end = _v22_previous_year(start), _v22_previous_year(end)
    current = _v20_period_count(sales, start, end)
    previous = _v20_period_count(sales, old_start, old_end)
    return {"now": current, "before": previous, "pct": _v22_pct(current, previous),
            "range": f"{old_start:%Y.%m.%d}~{old_end:%m.%d}"}


def _v22_dates_for_month(month, as_of):
    """선택 월의 전년 비교는 오늘이 포함된 미완료 월만 어제까지 자른다."""
    year, month_no = map(int, month.split("-"))
    start = date(year, month_no, 1)
    if month_no == 12:
        next_start = date(year + 1, 1, 1)
    else:
        next_start = date(year, month_no + 1, 1)
    end = min(next_start - timedelta(days=1), as_of)
    return start, end


@st.cache_data(ttl=180, max_entries=4, show_spinner=False)
def _v22_repurchase(sales, as_of):
    """60일 비율 외에 30/90일 코호트 비교. 이탈 위험 목록은 생성하지 않음."""
    out = {"median": None, "intervals": [], "cohorts": {30: (0, 0), 60: (0, 0), 90: (0, 0)}}
    if sales.empty or "analysis_customer_key" not in sales:
        return out
    valid = sales.loc[
        (sales["analysis_customer_key"].astype(str).str.strip() != "")
        & (sales["local_day"] <= as_of),
        ["analysis_customer_key", "local_day"]
    ].drop_duplicates().sort_values(["analysis_customer_key", "local_day"])
    if valid.empty:
        return out
    eligible = {window: 0 for window in (30, 60, 90)}
    converted = {window: 0 for window in (30, 60, 90)}
    intervals = []
    for _, group in valid.groupby("analysis_customer_key", sort=False):
        dates = group["local_day"].tolist()
        first = dates[0]
        second_gap = (dates[1] - first).days if len(dates) >= 2 else None
        for window in (30, 60, 90):
            if (as_of - first).days >= window:
                eligible[window] += 1
                if second_gap is not None and second_gap <= window:
                    converted[window] += 1
        intervals.extend([(b - a).days for a, b in zip(dates, dates[1:])])
    out["intervals"] = intervals
    out["median"] = float(pd.Series(intervals).median()) if intervals else None
    out["cohorts"] = {w: (eligible[w], converted[w]) for w in (30, 60, 90)}
    return out


def _v22_customer_split(sales, start, end):
    """선택 기간에 주문한 고객 중 첫 구매일이 기간 안인 신규/이전 구매자 구분."""
    if sales.empty or start > end:
        return {"new": 0, "returning": 0, "unknown": 0}
    keyed = sales[sales["analysis_customer_key"].astype(str).str.strip() != ""]
    if keyed.empty:
        return {"new": 0, "returning": 0, "unknown": _v20_period_count(sales, start, end)}
    firsts = keyed.groupby("analysis_customer_key")["local_day"].min()
    active = keyed[keyed["local_day"].between(start, end)]["analysis_customer_key"].drop_duplicates()
    first_dates = active.map(firsts)
    new = int(first_dates.ge(start).sum())
    returning = int(first_dates.lt(start).sum())
    unknown = int(sales[sales["local_day"].between(start, end)]["analysis_customer_key"].astype(str).str.strip().eq("").sum())
    return {"new": new, "returning": returning, "unknown": unknown}


def _v22_channel_compare(sales, start, end):
    """채널별 현재 vs 전년 동기간 (주문 건수)."""
    rows=[]
    for channel in ("쿠팡", "네이버"):
        data = sales[sales["channel"] == channel]
        r = _v22_yoy(data, start, end)
        rows.append({"채널": channel, "현재 주문": r["now"] or 0,
                     "작년 동기간": r["before"] or 0, "전년 대비": _v22_pct_label(r["pct"])})
    return pd.DataFrame(rows)


def _v22_month_frame(sales, as_of):
    """자료가 있는 전체 기간을 달력 월로 복원. 진행 중인 이번 달은 진행 중으로 표시."""
    if sales.empty:
        return pd.DataFrame(columns=["month", "orders", "short", "state"])
    month_counts = sales[sales["local_day"] <= as_of].groupby("month").size()
    if month_counts.empty:
        return pd.DataFrame(columns=["month", "orders", "short", "state"])
    month_index = pd.period_range(month_counts.index.min(),
                                  min(month_counts.index.max(),as_of.strftime("%Y-%m")),freq="M")
    frame = pd.DataFrame({"month": month_index.astype(str)})
    frame["orders"] = frame["month"].map(month_counts).fillna(0).astype(int)
    frame["short"] = frame["month"].str[2:4] + "." + frame["month"].str[5:7]
    frame["state"] = frame["month"].apply(lambda x: "진행 중" if x == as_of.strftime("%Y-%m") else "완료")
    return frame


def _v22_kpi(title, current, compare, note, suffix="건"):
    if current is None:
        st.metric(title, "—", delta=None)
    else:
        st.metric(title, f"{current:,}{suffix}",
                  delta=(f"전년 {_v22_pct_label(compare['pct'])}" if compare["pct"] is not None else None))
    st.caption(note)


def _v22_brief(yoy_current, yoy_last, yoy_year, seven_delta, last_db, as_of):
    """AI 호출 없이 확정 DB 숫자만 이용한 간결한 브리핑."""
    parts=[]
    if yoy_current["pct"] is not None:
        word="증가" if yoy_current["pct"] >= 0 else "감소"
        parts.append(f"이번 달 주문이 전년 동기간 대비 {abs(yoy_current['pct']):.1f}% {word}했습니다.")
    else:
        parts.append("이번 달은 전년 동기간에 주문 기록이 없어 증가율을 계산할 수 없습니다.")
    if seven_delta is not None:
        word="증가" if seven_delta >= 0 else "감소"
        parts.append(f"최근 완료 7일 주문은 직전 7일보다 {abs(seven_delta):.1f}% {word}했습니다.")
    if last_db is None:
        parts.append("주문 DB에 날짜 기록이 없습니다.")
    elif (as_of - last_db).days >= 3:
        parts.append(f"마지막 주문 기록은 {last_db:%m/%d}입니다. 최근 엑셀이 아직 저장되지 않았다면 비교값이 작게 나올 수 있습니다.")
    return parts


def _v22_monthly_selection_chart(visible, as_of):
    """12~24개월 막대 선택 시 고객이 수치를 크게 읽을 수 있게 반환."""
    selection = alt.selection_point(name="month_pick_v22", fields=["month"], on="click", clear="dblclick")
    xorder = list(visible["month"])
    chart = alt.Chart(visible).mark_bar(cornerRadiusTopLeft=5, cornerRadiusTopRight=5).encode(
        x=alt.X("month:N", sort=xorder, title=None,
                axis=alt.Axis(labelAngle=0, labelFontSize=15, labelPadding=12,
                              labelOverlap="greedy", labelExpr="substring(datum.label,2,4) + '.' + substring(datum.label,5,7)")),
        y=alt.Y("orders:Q", title="주문 건수", axis=alt.Axis(tickMinStep=1)),
        color=alt.condition(selection, alt.value("#1d4ed8"), alt.value("#7aa7ee")),
        tooltip=[alt.Tooltip("month:N", title="연·월"), alt.Tooltip("orders:Q",title="주문 건수",format=",d"),
                 alt.Tooltip("3개월 평균:Q",title="3개월 평균", format=".1f"),
                 alt.Tooltip("state:N",title="상태")]
    ).add_params(selection)
    state = _v20_big_chart(chart, height=350, key="v22_month_selection", selectable=True)
    chosen = None
    try:
        selected = state.selection.get("month_pick_v22", [])
        if selected and isinstance(selected, list):
            chosen = selected[0].get("month")
    except (AttributeError, KeyError, TypeError, IndexError):
        pass
    if chosen not in set(visible["month"]):
        chosen = visible["month"].iloc[-1]
    return chosen


def _v22_actual_weekday(inventory_history, as_of):
    """하루 전체 다섯 제품이 저장된 날짜만 실제 출고 요일 평균에 반영."""
    calendar = _recorded_inventory_calendar(inventory_history)
    start = as_of - timedelta(weeks=8)
    byday = {i: [] for i in range(7)}
    for d, inventory_items in calendar.items():
        if not (start <= d <= as_of) or d.weekday() >= 5:
            continue
        if all(key in inventory_items for key in KEYS):
            byday[d.weekday()].append(sum(_safe_int(inventory_items[key]) for key in KEYS))
    names = ["월","화","수","목","금"]
    return pd.DataFrame({"요일": names, "평균 출고(병)": [round(sum(byday[i])/len(byday[i]),1) if byday[i] else None for i in range(5)],
                         "기록일": [len(byday[i]) for i in range(5)]})


if tab_dashboard.open:
    with tab_dashboard:
        today_v22 = datetime.now(KST).date()
        as_of = today_v22 - timedelta(days=1)  # 완료된 날짜까지만 비교
        sales = _v20_orders(orders_db)
        sales = sales[sales["local_day"] <= as_of].copy() if not sales.empty else sales
        customer_total = len(customer_df)
        repeat_total = int((customer_df["order_count"] >= 2).sum()) if not customer_df.empty else 0
        vip_total = int((customer_df["order_count"] >= 5).sum()) if not customer_df.empty else 0

        this_month_start = today_v22.replace(day=1)
        last_month_end = this_month_start - timedelta(days=1)
        last_month_start = last_month_end.replace(day=1)
        this_year_start = today_v22.replace(month=1, day=1)
        yoy_current = _v22_yoy(sales, this_month_start, as_of)
        yoy_last = _v22_yoy(sales, last_month_start, last_month_end)
        yoy_ytd = _v22_yoy(sales, this_year_start, as_of)
        week_now, week_prev, week_change = _v20_compare(sales, 7, today_v22)
        thirty_now, thirty_prev, thirty_change = _v20_compare(sales, 30, today_v22)
        latest_record = max(sales["local_day"]) if not sales.empty else None

        st.title("📊 식혜명가 경영 대시보드 V22")
        st.caption(f"2026년형 판매·고객·생산 관리 · 비교 마감 {as_of:%Y.%m.%d} (어제까지) · **주문 건수 기준**")

        a,b,c = st.columns(3)
        with a:
            _v22_kpi("이번 달 주문", yoy_current["now"], yoy_current,
                     f"{this_month_start:%m/%d}~{as_of:%m/%d}, 작년 같은 날짜까지")
        with b:
            _v22_kpi("지난달 확정 주문", yoy_last["now"], yoy_last,
                     f"{last_month_start:%Y.%m} 전체 vs 전년 동월")
        with c:
            _v22_kpi("올해 누적 주문", yoy_ytd["now"], yoy_ytd,
                     f"1/1~{as_of:%m/%d}, 작년 동기간")

        a,b,c = st.columns(3)
        a.metric("최근 완료 7일", f"{week_now:,}건", delta=_v20_delta(week_change),
                 delta_color="off" if week_change is None else "normal")
        b.metric("최근 완료 30일", f"{thirty_now:,}건", delta=_v20_delta(thirty_change),
                 delta_color="off" if thirty_change is None else "normal")
        c.metric("재구매 고객 비중", f"{repeat_total/customer_total*100:.1f}%" if customer_total else "—",
                 help="DB 전체 고객 중 주문 2회 이상 고객 비중. 기간 코호트 재구매율과는 다릅니다.")

        with st.expander("📌 전년 비교 상세 수치 확인", expanded=False):
            data=[]
            for label,item in [("이번 달 (완료일 기준)",yoy_current),("지난달 전체",yoy_last),("올해 누적",yoy_ytd)]:
                data.append({"구분":label,"올해":item["now"] if item["now"] is not None else "—",
                             "작년":item["before"] if item["before"] is not None else "—",
                             "증감률":_v22_pct_label(item["pct"]),"작년 비교기간":item["range"] or "—"})
            st.dataframe(pd.DataFrame(data), use_container_width=True, hide_index=True)
            st.caption("작년 0건이면 증가율은 계산하지 않습니다. 2월 29일은 전년 2월 28일과 비교합니다. 주문 취소·환불은 원본 DB에 반영된 범위에서만 집계됩니다.")

        st.subheader("📝 이번 주 판매 브리핑")
        for i, line in enumerate(_v22_brief(yoy_current,yoy_last,yoy_ytd,week_change,latest_record,as_of), start=1):
            st.markdown(f"**{i}.** {line}")
        if latest_record and (as_of-latest_record).days >= 3:
            st.warning(f"주문 DB 마지막 기록 {latest_record:%Y.%m.%d}. 최신 파일을 저장하지 않았다면 비교 수치가 실제보다 적을 수 있어요.")
        if not sales.empty and sales["local_day"].min() > _v22_previous_year(this_month_start):
            st.info("전년도 해당 기간의 DB 기록이 충분하지 않을 수 있습니다. 집계 시작일을 확인하세요.")

        st.divider()
        panels = st.tabs(["📈 판매·전년 비교", "👥 신규·재구매", "🏭 생산·재고", "🛒 채널·명절"],
                         on_change="rerun", key="crm_dashboard_panels_v22")

        if panels[0].open:
            with panels[0]:
                st.subheader("월별 주문 흐름")
                view = st.segmented_control("표시 기간", ["최근 12개월", "최근 24개월", "연도별"],
                                             default="최근 12개월",key="v22_month_period")
                monthly = _v22_month_frame(sales,as_of)
                if monthly.empty:
                    st.info("표시할 주문 데이터가 없습니다.")
                else:
                    if view == "최근 12개월":
                        visible=monthly.tail(12).copy()
                    elif view == "최근 24개월":
                        visible=monthly.tail(24).copy()
                    else:
                        years=sorted(monthly["month"].str[:4].unique().tolist(),reverse=True)
                        yr=st.selectbox("연도 선택",years,key="v22_year")
                        visible=monthly[monthly["month"].str.startswith(yr)].copy()
                    visible["3개월 평균"] = visible["orders"].rolling(3,min_periods=1).mean().round(1)
                    selected_month = _v22_monthly_selection_chart(visible,as_of)
                    month_count=int(visible.loc[visible["month"]==selected_month,"orders"].iloc[0])
                    s,e=_v22_dates_for_month(selected_month,as_of)
                    comparison=_v22_yoy(sales,s,e)
                    p,q=st.columns(2)
                    p.metric(f"선택 월 {selected_month}",f"{month_count:,}건")
                    q.metric("선택 월 전년 대비",_v22_pct_label(comparison["pct"]),
                             help="진행 중인 달은 같은 날짜까지만 비교하며, 전년도 0건이면 비교 불가입니다.")
                    st.caption("막대를 클릭해 수치를 크게 확인하세요. 연도를 포함한 축 글씨는 가로 방향 고정입니다. 이번 달은 진행 중인 달이에요.")
                    with st.expander("월별 수치 표 자세히 보기"):
                        st.dataframe(visible[["month","orders","3개월 평균","state"]].rename(columns={"month":"월","orders":"주문 건수","state":"기록 상태"}),
                                     use_container_width=True,hide_index=True)

                st.divider()
                st.subheader("요일별 주문 패턴")
                anchor=today_v22-timedelta(days=today_v22.weekday())
                start=anchor-timedelta(weeks=8)
                prior=sales[sales["local_day"].between(start,anchor-timedelta(days=1))]
                order_days=prior.groupby("weekday").size().to_dict()
                names=["월","화","수","목","금","토","일"]
                pattern=pd.DataFrame({"요일":names,"하루 평균 주문": [round(order_days.get(i,0)/8,1) for i in range(7)]})
                daychart=alt.Chart(pattern).mark_bar(size=46,cornerRadiusTopLeft=5,cornerRadiusTopRight=5).encode(
                    x=alt.X("요일:N",sort=names,title=None,
                            axis=alt.Axis(labelAngle=0,labelFontSize=18,labelPadding=12,labelOverlap=False)),
                    y=alt.Y("하루 평균 주문:Q",title="일 평균 주문(건)"),
                    tooltip=[alt.Tooltip("요일:N"),alt.Tooltip("하루 평균 주문:Q",format=".1f")],
                    color=alt.value("#2563eb"))
                values=alt.Chart(pattern).mark_text(dy=-11,fontSize=17,fontWeight="bold",color="#1f2937").encode(
                    x=alt.X("요일:N",sort=names),y="하루 평균 주문:Q",text=alt.Text("하루 평균 주문:Q",format=".1f"))
                _v20_big_chart(daychart+values,320)
                st.caption(f"직전 완료된 8주({start:%Y-%m-%d}~{(anchor-timedelta(days=1)):%Y-%m-%d})의 **주문일 기준** 요일당 평균. 실제 출고요일은 생산·재고 화면에서 별도 확인합니다.")

        if panels[1].open:
            with panels[1]:
                st.subheader("신규 고객 vs 기존 고객")
                split = _v22_customer_split(sales,this_month_start,as_of)
                ca,cb,cc=st.columns(3)
                ca.metric("이번 달 신규 구매 고객",f"{split['new']:,}명")
                cb.metric("이번 달 기존 구매 고객",f"{split['returning']:,}명")
                total_known=split['new']+split['returning']
                cc.metric("기존 고객 비중",f"{split['returning']/total_known*100:.1f}%" if total_known else "—")
                st.caption("이 달에 구매한 고유 고객을 최초 구매일 기준으로 구분합니다. DB에 없는 과거 주문은 반영할 수 없습니다.")
                if split['unknown']:
                    st.warning(f"고객 식별자가 없는 주문 {split['unknown']:,}건은 신규/기존 분류에서 제외했습니다.")
                st.divider()
                st.subheader("첫 구매 후 재구매 전환율")
                report=_v22_repurchase(sales,as_of)
                cols=st.columns(3)
                for col,days in zip(cols,(30,60,90)):
                    eligible,converted=report['cohorts'][days]
                    with col:
                        st.metric(f"{days}일 내 재구매", f"{converted/eligible*100:.1f}%" if eligible else "—",
                                  help=f"첫 구매 후 {days}일의 관찰 기간이 지난 고객만 분모에 포함")
                        st.caption(f"{converted:,}명 / 대상 {eligible:,}명")
                st.caption("같은 날 여러 주문을 해도 '두 번째 구매일'로 세지 않습니다. 기간별로 분석 대상 고객 수가 다릅니다.")
                st.metric("재구매 간격 중앙값",f"{report['median']:.0f}일" if report['median'] is not None else "기록 부족")
                intervals=report['intervals']
                if intervals:
                    bins=[("1~7일",1,7),("8~30일",8,30),("31~60일",31,60),("61~90일",61,90),("91일 이상",91,100000)]
                    frame=pd.DataFrame({"재구매 간격":[b[0] for b in bins],"건수":[sum(b[1]<=g<=b[2] for g in intervals) for b in bins]})
                    bars=alt.Chart(frame).mark_bar(size=30,cornerRadiusEnd=4).encode(
                        y=alt.Y("재구매 간격:N",sort=[b[0] for b in bins],title=None,axis=alt.Axis(labelAngle=0)),
                        x=alt.X("건수:Q",title="구매 간격 사례"),color=alt.value("#2563eb"),
                        tooltip=["재구매 간격:N",alt.Tooltip("건수:Q",format=",d")])
                    _v20_big_chart(bars,290)
                st.divider()
                st.subheader("고객 구매횟수 분포")
                counts=pd.to_numeric(customer_df["order_count"],errors="coerce").fillna(0)
                buck=[("1회",int((counts==1).sum())),("2회",int((counts==2).sum())),
                      ("3회",int((counts==3).sum())),("4회",int((counts==4).sum())),
                      ("5회 이상",int((counts>=5).sum()))]
                distribution=pd.DataFrame(buck,columns=["구매횟수","고객 수"])
                c=alt.Chart(distribution).mark_bar(size=30).encode(
                    y=alt.Y("구매횟수:N",sort=[x[0] for x in buck],title=None,axis=alt.Axis(labelAngle=0)),
                    x=alt.X("고객 수:Q",title="고객 수"),color=alt.value("#e59c29"),
                    tooltip=["구매횟수:N",alt.Tooltip("고객 수:Q",format=",d")])
                _v20_big_chart(c,290)

        if panels[2].open:
            with panels[2]:
                st.subheader("주간 생산량·재고 소진")
                st.caption("실제 병 수는 **재고관리에서 매일 최종 저장한 품목별 출고량**만 사용합니다. 기존 주문 건수로 생산량을 추측하지 않습니다.")
                history,inventory_error=fetch_inventory_rows()
                if inventory_error:
                    st.warning(inventory_error)
                else:
                    production_df,full_weeks=_v20_production(history,today_v22)
                    st.metric("최근 4주 중 평일 5일 모두 저장한 주",f"{full_weeks}/4주")
                    st.dataframe(production_df,use_container_width=True,hide_index=True)
                    st.caption("생산 참고량 = 주간 예상 출고 + 품목별 최소 확보재고 − 최신 재고. 출고 완전 기록 2주 이상, 7일 이내 재고 갱신시에만 표시해요.")
                    if full_weeks < 2:
                        st.info("완전히 기록한 평일 주가 2주 이상 쌓이면 예상 출고와 생산 참고량을 계산합니다. 미기록은 0출고로 계산하지 않습니다.")
                    else:
                        chart_data=production_df[pd.to_numeric(production_df['소진 예상(영업일)'],errors='coerce').notna()].copy()
                        if not chart_data.empty:
                            st.subheader("품목별 예상 재고 소진일")
                            chart_data['예상 영업일']=pd.to_numeric(chart_data['소진 예상(영업일)'])
                            c=alt.Chart(chart_data).mark_bar(size=25).encode(
                                y=alt.Y('품목:N',sort=[p[1] for p in ITEMS],title=None,axis=alt.Axis(labelAngle=0)),
                                x=alt.X('예상 영업일:Q',title='평균 판매 속도 기준 영업일'),
                                tooltip=['품목:N',alt.Tooltip('예상 영업일:Q',format='.1f')],color=alt.value('#f59e0b'))
                            _v20_big_chart(c,290)
                    st.divider()
                    st.subheader("실제 출고요일 패턴 (월~금)")
                    shipping_days = _v22_actual_weekday(history,today_v22)
                    if shipping_days['기록일'].sum() == 0:
                        st.info("아직 완전히 기록한 평일 출고 데이터가 없습니다.")
                    else:
                        st.dataframe(shipping_days,use_container_width=True,hide_index=True)
                        filled=shipping_days.dropna(subset=['평균 출고(병)']).copy()
                        if not filled.empty:
                            x=alt.Chart(filled).mark_bar(size=42).encode(
                                x=alt.X('요일:N',sort=['월','화','수','목','금'],title=None,
                                        axis=alt.Axis(labelAngle=0,labelFontSize=17)),
                                y=alt.Y('평균 출고(병):Q',title='병'),color=alt.value('#16a34a'),
                                tooltip=['요일:N',alt.Tooltip('평균 출고(병):Q',format='.1f')])
                            _v20_big_chart(x,260)
                        st.caption('실제 출고량은 5품목을 모두 저장한 날짜만 이용합니다. 요일별 표본 수가 다르므로 참고용으로 보세요.')

        if panels[3].open:
            with panels[3]:
                st.subheader("쿠팡 vs 네이버")
                channel_frame = _v22_channel_compare(sales,this_month_start,as_of)
                st.dataframe(channel_frame,use_container_width=True,hide_index=True)
                st.caption("이번 달 주문을 작년 같은 날짜까지 비교했습니다. 채널이 확인되지 않는 주문은 표에 포함하지 않았습니다.")
                periods=pd.period_range(end=pd.Period(as_of.strftime('%Y-%m'),freq='M'),periods=6,freq='M').astype(str).tolist()
                filtered=sales[sales['month'].isin(periods) & sales['channel'].isin(['쿠팡','네이버'])]
                grouped=filtered.groupby(['month','channel']).size().rename('주문 건수').reset_index()
                grid=pd.MultiIndex.from_product([periods,['쿠팡','네이버']],names=['month','channel']).to_frame(index=False)
                grouped=grid.merge(grouped,on=['month','channel'],how='left').fillna({'주문 건수':0})
                chart=alt.Chart(grouped).mark_bar().encode(
                    x=alt.X('month:N',sort=periods,title=None,
                            axis=alt.Axis(labelAngle=0,labelFontSize=15,
                                          labelExpr="substring(datum.label,2,4) + '.' + substring(datum.label,5,7)")),
                    y=alt.Y('주문 건수:Q',title='건'),
                    color=alt.Color('channel:N',title='판매채널',scale=alt.Scale(domain=['쿠팡','네이버'],range=['#2563eb','#16a34a'])),
                    xOffset='channel:N',tooltip=['month:N','channel:N',alt.Tooltip('주문 건수:Q',format=',.0f')])
                _v20_big_chart(chart,330)
                others=int((sales['channel']=='기타').sum()) if not sales.empty else 0
                if others:
                    st.warning(f"채널 분류 불명 주문 {others:,}건은 쿠팡/네이버 비교에서 제외했습니다.")

                st.divider()
                st.subheader("명절 주문 수요")
                holiday=st.segmented_control('명절 종류',['추석','설날'],default='추석',key='v22_holiday_kind')
                preset={'추석':(date(2026,9,25),date(2025,10,6)),
                        '설날':(date(2026,2,17),date(2025,1,29))}
                year_now,year_before=preset[holiday]
                c1,c2=st.columns(2)
                with c1:
                    event_now=st.date_input('이번 비교 명절 당일',value=year_now,key=f'v22_event_now_{holiday}')
                with c2:
                    event_old=st.date_input('이전 비교 명절 당일',value=year_before,key=f'v22_event_old_{holiday}')
                h=_v20_holiday(sales,event_now,event_old,today_v22)
                st.dataframe(h,use_container_width=True,hide_index=True)
                st.caption('명절 당일 기준 21~1일 전의 주문 건수를 비교합니다. 매년 명절 날짜가 바뀌므로 행사 날짜에 맞춰 분석하고 병 수와 구분합니다.')

        with st.expander('🤖 CRM AI 요약 (선택 실행)', expanded=False):
            ready=ai_status_line()
            if st.button('✨ 현재 CRM 자동 분석',type='primary',key='dashboard_ai_run_v22',disabled=not ready):
                with st.spinner('고객 개인정보를 제외한 집계 수치 분석 중...'):
                    st.session_state['dashboard_ai_result_v22']=auto_ai_crm_analysis(customer_df,orders_db)
            if st.session_state.get('dashboard_ai_result_v22'):
                st.markdown(st.session_state['dashboard_ai_result_v22'])

if tab_customers.open:
    with tab_customers:
        st.subheader("👤 고객 CRM")
        st.caption("이름/전화번호/주소/상품명으로 검색하고, VIP·재구매·블랙 고객을 색상으로 확인합니다.")

        search = st.text_input("통합 검색", placeholder="예: 고객명, 전화번호, 주소, 상품명")
        grade_options = sorted(customer_df["grade"].dropna().unique())
        selected_grades = st.multiselect("고객등급", grade_options, default=grade_options)

        black_options = sorted(customer_df["black_status"].dropna().unique())
        selected_black = st.multiselect("블랙상태", black_options, default=black_options)

        view = customer_df[
            (customer_df["grade"].isin(selected_grades)) &
            (customer_df["black_status"].isin(selected_black))
        ].copy()

        if search:
            search_mask = (
                view["customer_name"].astype(str).str.contains(search, case=False, na=False, regex=False)
                | view["phone"].astype(str).str.contains(search, case=False, na=False, regex=False)
                | view["address"].astype(str).str.contains(search, case=False, na=False, regex=False)
                | view["products"].astype(str).str.contains(search, case=False, na=False, regex=False)
            )
            view = view[search_mask]

        sort_col = st.selectbox("정렬", ["order_count", "total_amount", "last_order_date", "avg_order_amount"], index=0)
        view = view.sort_values(sort_col, ascending=False)

        show_cols = [
            "customer_name", "phone", "address", "grade", "black_status",
            "order_count", "total_amount", "avg_order_amount",
            "first_order_date", "last_order_date", "channels", "products"
        ]
        show_cols = [c for c in show_cols if c in view.columns]

        st.dataframe(style_customer_crm_table(view[show_cols]), use_container_width=True, hide_index=True, height=520)

        st.markdown("### 고객 CRM 구매이력 바로 확인")
        if view.empty:
            st.info("검색 결과가 없습니다.")
        else:
            history_options = view.copy()
            history_options["history_label"] = (
                history_options["customer_name"].astype(str)
                + " | "
                + history_options["phone"].astype(str)
                + " | "
                + history_options["order_count"].astype(str)
                + "회"
            )

            selected_history_label = st.selectbox("구매이력 확인할 고객 선택", history_options["history_label"].tolist())
            selected_history = history_options[history_options["history_label"] == selected_history_label].iloc[0]
            selected_key = selected_history["analysis_customer_key"]

            h1, h2, h3, h4 = st.columns(4)
            h1.metric("주문횟수", f"{int(selected_history['order_count']):,}회")
            h2.metric("총 주문금액", f"{int(selected_history.get('total_amount', 0)):,}원")
            h3.metric("건당 평균금액", f"{int(selected_history.get('avg_order_amount', 0)):,}원")
            h4.metric("최근 주문일", str(pd.to_datetime(selected_history["last_order_date"]).date()))

            notes_df = fetch_customer_notes()
            memo_text = ""
            if not notes_df.empty and selected_key in notes_df["customer_key"].astype(str).values:
                memo_text = notes_df[notes_df["customer_key"].astype(str) == str(selected_key)]["memo"].iloc[0]
            if memo_text:
                st.markdown(f"**고객 메모**  \n{memo_text}")

            history_df = orders_db[orders_db["analysis_customer_key"] == selected_key].sort_values("order_date", ascending=False)
            st.dataframe(history_df, use_container_width=True, hide_index=True, height=300)



if tab_detail.open:
    with tab_detail:
        st.subheader("🔎 고객 상세")
        st.caption("고객을 선택하면 주문횟수, 첫/최근 주문일, 전체 주문내역, 고객 메모를 볼 수 있습니다.")

        customer_view = customer_df.sort_values(["order_count", "last_order_date"], ascending=[False, False]).copy()
        customer_view["label"] = (
            customer_view["customer_name"].astype(str)
            + " | "
            + customer_view["phone"].astype(str)
            + " | "
            + customer_view["order_count"].astype(str)
            + "회"
        )

        keyword = st.text_input("이름/전화번호/주소 검색")
        if keyword:
            mask = (
                customer_view["customer_name"].astype(str).str.contains(keyword, case=False, na=False, regex=False)
                | customer_view["phone"].astype(str).str.contains(keyword, case=False, na=False, regex=False)
                | customer_view["address"].astype(str).str.contains(keyword, case=False, na=False, regex=False)
            )
            customer_view = customer_view[mask]

        if customer_view.empty:
            st.info("검색 결과가 없습니다.")
        else:
            selected_label = st.selectbox("고객 선택", customer_view["label"].tolist())
            selected = customer_view[customer_view["label"] == selected_label].iloc[0]
            key = selected["analysis_customer_key"]

            c1, c2, c3, c4, c5, c6 = st.columns(6)
            c1.metric("주문횟수", f"{int(selected['order_count']):,}회")
            c2.metric("등급", str(selected["grade"]))
            c3.metric("총 주문금액", f"{int(selected.get('total_amount', 0)):,}원")
            c4.metric("건당 평균금액", f"{int(selected.get('avg_order_amount', 0)):,}원")
            c5.metric("첫 주문일", str(pd.to_datetime(selected["first_order_date"]).date()))
            c6.metric("최근 주문일", str(pd.to_datetime(selected["last_order_date"]).date()))

            st.write("고객정보")
            info_df = pd.DataFrame([{
                "고객명": selected["customer_name"],
                "전화번호": selected["phone"],
                "주소": selected["address"],
                "이용채널": selected["channels"],
                "구매상품": selected["products"],
            }])
            st.dataframe(info_df, use_container_width=True, hide_index=True)

            st.write("고객 메모 / 태그")
            notes_df = fetch_customer_notes()
            old_memo = ""
            if not notes_df.empty and key in notes_df["customer_key"].astype(str).values:
                old_memo = notes_df[notes_df["customer_key"].astype(str) == str(key)]["memo"].iloc[0]

            tag_options = [
                "단호박매니아", "식혜매니아", "혼합구매", "대량주문", "소량주문",
                "신규", "재구매", "VIP", "VVIP",
                "쿠폰반응", "행사반응", "선물용", "재구매유도",
                "CS주의", "배송주의", "블랙",
                "회사구매", "학교구매", "단체구매"
            ]
            selected_tags = st.multiselect("고객 태그", tag_options)
            memo_input = st.text_area("메모 입력", value=old_memo, placeholder="예: 단호박 선호 / 행사 대량주문 / CS 주의")
            save_text = "[태그] " + ", ".join(selected_tags) + "\n" + memo_input if selected_tags else memo_input

            if st.button("고객 메모 저장", type="primary"):
                save_customer_note(key, selected["customer_name"], selected["phone"], selected["address"], save_text)
                st.success("고객 메모 저장 완료")

            st.write("전체 주문내역")
            history = orders_db[orders_db["analysis_customer_key"] == key].sort_values("order_date", ascending=False)
            st.dataframe(history, use_container_width=True, hide_index=True)


if tab_vip.open:
    with tab_vip:
        st.subheader("⭐ VIP 고객")
        st.caption("기본 5회 이상, 필요하면 기준을 바꿔 확인할 수 있습니다.")

        vip_min = st.slider("VIP 기준 주문횟수", min_value=2, max_value=20, value=5, step=1)
        vip_df = customer_df[customer_df["order_count"] >= vip_min].sort_values(["order_count", "last_order_date"], ascending=[False, False])

        a, b = st.columns(2)
        a.metric("VIP 고객수", f"{len(vip_df):,}명")
        b.metric("최고 주문횟수", f"{int(vip_df['order_count'].max()) if not vip_df.empty else 0:,}회")

        st.dataframe(vip_df, use_container_width=True, hide_index=True)



if tab_ai.open:
    with tab_ai:
        st.subheader("🤖 AI CRM 비서")
        st.caption("고객 개인정보는 보내지 않고 누적 주문·재구매율·구매횟수 등 집계만 전송합니다.")
        ai_ready = ai_status_line()
        if not ai_ready:
            st.code('AI_PROVIDER = "gemini"\nGEMINI_API_KEY = "실제 API 키"\nGEMINI_MODEL = "gemini-3.5-flash-lite"', language="toml")
            st.caption("설정 경로: Streamlit Cloud → 앱 → Settings → Secrets · 비밀키는 채팅에 보내지 마세요.")

        st.markdown("### 🚀 자동 CRM 분석")
        if st.button("✨ 자동 CRM 분석 실행", type="primary", use_container_width=True,
                     key="ai_tab_auto_v19", disabled=not ai_ready):
            with st.spinner("집계 데이터를 분석 중입니다..."):
                st.session_state["ai_tab_auto_result_v19"] = auto_ai_crm_analysis(customer_df, orders_db)
        if st.session_state.get("ai_tab_auto_result_v19"):
            st.markdown(st.session_state["ai_tab_auto_result_v19"])

        st.divider()
        st.markdown("### 💬 직접 질문")
        quick_q = st.selectbox("빠른 질문", [
            "현재 CRM 상태를 짧게 분석해줘",
            "재구매율을 높이려면 뭘 하면 좋을까?",
            "VIP 고객 관리 전략을 추천해줘",
            "오늘 주문 분석 결과에서 주의할 점을 알려줘",
            "마케팅 문자 보낼 고객 기준을 추천해줘",
            "2회 구매 고객을 3회 구매로 전환하는 방법을 알려줘",
            "VIP 고객에게 보낼 문자 문구를 작성해줘",
        ])
        custom_q = st.text_area("직접 질문", placeholder="예: 3회 이상 구매고객에게 어떤 이벤트를 하면 좋을까?")
        question = custom_q.strip() if custom_q.strip() else quick_q
        if st.button("AI에게 질문하기", use_container_width=True,
                     key="ai_tab_question_v19", disabled=not ai_ready):
            with st.spinner("분석 중입니다..."):
                st.session_state["ai_tab_question_result_v19"] = ask_ai_crm(question, customer_df, orders_db)
        if st.session_state.get("ai_tab_question_result_v19"):
            st.markdown(st.session_state["ai_tab_question_result_v19"])

if tab_black.open:
    with tab_black:
        st.subheader("블랙리스트 / 주의고객 관리")

        customer_options = customer_df.sort_values(["order_count", "last_order_date"], ascending=[False, False]).copy()
        customer_options["label"] = (
            customer_options["customer_name"].astype(str)
            + " | "
            + customer_options["phone"].astype(str)
            + " | "
            + customer_options["address"].astype(str).str[:40]
        )

        selected_label = st.selectbox("고객 선택", customer_options["label"].tolist())
        selected = customer_options[customer_options["label"] == selected_label].iloc[0]

        with st.form("black_form"):
            status = st.selectbox("상태", ["정상", "주의", "블랙"], index=0)
            reason = st.selectbox("사유", ["", "반복 환불 요구", "허위 클레임", "배송 트집", "폭언/욕설", "고의 반품", "기타"])
            incident_date = st.date_input("발생일", value=date.today())
            result = st.text_input("처리 결과", placeholder="예: 재배송함 / 환불함 / 다음 주문 시 확인")
            memo = st.text_area("상세 기록", placeholder="어떤 문제가 있었는지 기록")
            submitted = st.form_submit_button("저장")

        if submitted:
            row = {
                "customer_key": str(selected["analysis_customer_key"]),
                "customer_name": str(selected["customer_name"]),
                "phone": str(selected["phone"]),
                "address": str(selected["address"]),
                "status": status,
                "reason": reason,
                "memo": memo,
                "incident_date": str(incident_date),
                "result": result,
            }
            save_blacklist(row)
            st.success("블랙리스트 기록 저장 완료. 새로고침하면 반영됩니다.")

        st.subheader("현재 블랙리스트")
        st.dataframe(black_df.sort_values("updated_at", ascending=False) if "updated_at" in black_df.columns else black_df, use_container_width=True, hide_index=True)

        if not black_df.empty:
            delete_key = st.selectbox("삭제할 기록 선택", black_df["customer_key"].tolist())
            if st.button("선택 기록 삭제"):
                delete_blacklist(delete_key)
                st.success("삭제 완료. 새로고침하면 반영됩니다.")

if tab_orders.open:
    with tab_orders:
        st.subheader("저장된 주문 DB")
        st.dataframe(orders_db.sort_values("order_date", ascending=False), use_container_width=True, hide_index=True)

if tab_download.open:
    with tab_download:
        st.subheader("엑셀 다운로드")

        vip = customer_df[customer_df["order_count"] >= 5].copy()
        repeat = customer_df[customer_df["order_count"] >= 2].copy()
        black_customers = customer_df[customer_df["black_status"].isin(["주의", "블랙"])].copy()

        summary = pd.DataFrame({
            "항목": ["누적 주문", "누적 고객", "재구매 고객", "재구매율", "VIP", "주의/블랙"],
            "값": [len(orders_db), len(customer_df), len(repeat), f"{100 * len(repeat) / len(customer_df) if len(customer_df) else 0:.1f}%", len(vip), len(black_customers)]
        })

        download_sheets = {
            "요약": summary,
            "고객CRM": customer_df.sort_values(["order_count", "last_order_date"], ascending=[False, False]),
            "VIP": vip,
            "재구매": repeat,
            "블랙리스트": black_customers,
            "주문DB": orders_db.sort_values("order_date", ascending=False),
        }

        today_customer_analysis = st.session_state.get("today_customer_analysis", pd.DataFrame())
        today_detail_analysis = st.session_state.get("today_detail_analysis", pd.DataFrame())
        if today_customer_analysis is not None and not today_customer_analysis.empty:
            download_sheets["오늘주문분석"] = today_customer_analysis
        if today_detail_analysis is not None and not today_detail_analysis.empty:
            download_sheets["오늘주문상세"] = today_detail_analysis

        st.caption("⚡ 엑셀 파일은 요청할 때만 생성합니다. 다른 화면에서 이동할 때는 대용량 파일 생성 작업을 실행하지 않아요.")
        if st.button("📥 CRM 전체 엑셀 파일 만들기", type="primary", use_container_width=True, key="crm_excel_prepare_v21"):
            with st.spinner("고객·주문 기록을 엑셀로 정리하는 중입니다..."):
                st.session_state["_crm_excel_bytes_v21"] = to_excel_bytes(download_sheets)
                st.session_state["_crm_excel_name_v21"] = f"sikhye_saved_crm_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        if st.session_state.get("_crm_excel_bytes_v21"):
            st.download_button(
                "✅ 준비된 엑셀 다운로드",
                data=st.session_state["_crm_excel_bytes_v21"],
                file_name=st.session_state["_crm_excel_name_v21"],
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )